import sys
import subprocess
import os
import tempfile
import shutil

# ==========================================
# VERIFICAÇÃO E INSTALAÇÃO AUTOMÁTICA
# ==========================================
def checar_dependencias():
    dependencias = {
        "pandas": "pandas",
        "openpyxl": "openpyxl",
        "customtkinter": "customtkinter",
        "PIL": "Pillow",
        "win32com": "pywin32",
        "pypdf": "pypdf",
        "docx": "python-docx",
        "pypdfium2": "pypdfium2"
    }

    for modulo, pacote in dependencias.items():
        try:
            __import__(modulo)
        except ImportError:
            print(f"Instalando pacote necessário: {pacote}...")
            subprocess.check_call([sys.executable, "-m", "pip", "install", pacote])

checar_dependencias()

# ==========================================
# IMPORTS GERAIS
# ==========================================
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import datetime
from datetime import datetime as dt
import customtkinter as ctk
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from PIL import Image
import time
import pythoncom
import pywintypes
import traceback
import win32com.client as win32
import win32com.client.dynamic as dynamic
import pandas as pd
import threading
import queue
import re
import gc

# =====================================================================
# RESOLUÇÃO DINÂMICA DE RECURSOS (Compatível com PyInstaller)
# =====================================================================
def obter_caminho_recurso(nome_arquivo):
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        caminho_interno = os.path.join(sys._MEIPASS, nome_arquivo)
        if os.path.exists(caminho_interno):
            return caminho_interno

    diretorio_script = os.path.dirname(os.path.abspath(__file__))
    caminho_direto = os.path.join(diretorio_script, nome_arquivo)
    if os.path.exists(caminho_direto):
        return caminho_direto

    return nome_arquivo

CAMINHO_LOGO = obter_caminho_recurso("gree.png")
CAMINHO_ICO = obter_caminho_recurso("gree.ico")
CAMINHO_LOGO_COR = obter_caminho_recurso("logo gree colorida.png")
# =====================================================================

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

# ==========================================
# FUNÇÕES E VARIÁVEIS GLOBAIS - CRM
# ==========================================
LISTA_CORES = [
    "8EA9DB", "A9D08E", "F4B084", "FF99CC", "B4C6E7",
    "FFD966", "C6E0B4", "CC99FF", "E2EFDA", "99CCFF",
    "F8CBAD", "D9E1F2"
]

def converter_numero(valor):
    if valor is None or str(valor).strip() == "": return None
    try: return int(str(valor).strip().replace(".", "").replace(",", ""))
    except: return valor

def moeda_americana_para_brasileira(valor):
    if valor is None or valor == "": return None
    if isinstance(valor, (int, float)): return float(valor)
    texto = str(valor).strip().replace(",", "")
    try: return float(texto)
    except: return None

def converter_cubagem(valor):
    if valor is None or valor == "": return None
    if isinstance(valor, (int, float)): return float(valor)
    texto = str(valor).strip()
    if "." in texto and "," not in texto:
        texto = texto.replace(".", ",")
    texto = texto.replace(".", "").replace(",", ".")
    try: return float(texto)
    except: return None

def remover_sufixo(valor):
    if valor is None: return ""
    texto = str(valor).strip()
    if len(texto) >= 2 and (texto.endswith("-1") or texto.endswith("-2")):
        return texto[:-2]
    return texto

def obter_campo_dinamico(pt, nome_procurado):
    try:
        for i in range(1, pt.PivotFields().Count + 1):
            campo = pt.PivotFields(i)
            nome_campo = str(campo.Name).strip().lower()
            if nome_campo == nome_procurado.strip().lower():
                return campo
        for i in range(1, pt.PivotFields().Count + 1):
            campo = pt.PivotFields(i)
            nome_campo = str(campo.Name).strip().lower()
            if nome_procurado.strip().lower() in nome_campo:
                return campo
    except:
        pass
    try:
        return pt.PivotFields(nome_procurado)
    except:
        return None

def _codigo_com_error(e):
    if isinstance(e, pywintypes.com_error):
        args = e.args or ()
        return args[0] if args else None
    return None

def _deve_tentar_reabrir_excel(e):
    codigo = _codigo_com_error(e)
    return codigo in {-2146827284, -2147023174}

def _normalizar_titulo_planilha(titulo):
    texto = str(titulo or "").strip()
    if not texto:
        return "Planilha"
    return texto[:31]

def _fechar_excel_forcado(excel):
    if excel is None: return None
    try:
        for workbook in list(excel.Workbooks):
            try: workbook.Close(SaveChanges=False)
            except Exception: pass
    except Exception: pass
    try: excel.Quit()
    except Exception: pass
    return None

def _reinicializar_excel(excel):
    _fechar_excel_forcado(excel)
    try: pythoncom.CoUninitialize()
    except Exception: pass
    try: pythoncom.CoInitialize()
    except Exception: pass
    # Linhas removidas para não fechar as planilhas do usuário no Windows
    # try: subprocess.run(["taskkill", "/F", "/IM", "excel.exe"], check=False, capture_output=True, text=True, timeout=10)
    # except Exception: pass
    return abrir_excel_silencioso()

def abrir_excel_silencioso():
    excel = None
    try:
        pythoncom.CoInitialize()
    except Exception:
        pass

    try:
        caminho_cache = os.path.join(tempfile.gettempdir(), 'gen_py')
        if os.path.exists(caminho_cache):
            shutil.rmtree(caminho_cache, ignore_errors=True)
    except:
        pass

    for tentativa in range(4):
        try:
            # win32.DispatchEx força a abertura de uma nova instância em segundo plano
            excel = win32.DispatchEx("Excel.Application")
            _ = excel.Workbooks
            break
        except Exception:
            try:
                excel = win32.Dispatch("Excel.Application")
                _ = excel.Workbooks
                break
            except Exception:
                excel = None

        if excel is None:
            time.sleep(1)

    if excel is None:
        return None

    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        excel.AskToUpdateLinks = False
    except Exception:
        pass

    return excel

def abrir_workbook_com_retry(excel, caminho_absoluto, tentativas=8, atraso=1.0):
    instancia_excel = excel
    caminho_absoluto = os.path.abspath(caminho_absoluto)

    for tentativa in range(1, tentativas + 1):
        try:
            wb = instancia_excel.Workbooks.Open(
                Filename=caminho_absoluto,
                ReadOnly=False,
                UpdateLinks=False,
                Notify=False,
                IgnoreReadOnlyRecommended=True,
                Editable=True,
            )
            return instancia_excel, wb
        except pywintypes.com_error as e:
            codigo = _codigo_com_error(e)
            if codigo == -2147418111:
                time.sleep(atraso)
                continue
            elif _deve_tentar_reabrir_excel(e) and tentativa < tentativas:
                instancia_excel = _reinicializar_excel(instancia_excel)
                if instancia_excel is None: raise
                time.sleep(atraso)
                continue
            raise
        except Exception: raise
    return instancia_excel, None

def letter_to_index(letter):
    letter = letter.upper()
    result = 0
    for char in letter:
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result - 1

def safe_str(val):
    if pd.isna(val):
        return ''
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()

# ==========================================
# CÓDIGO 1: APLICATIVO CRM (Layout em 2 Colunas)
# ==========================================
class CRMFormatter(ctk.CTkFrame):
    def __init__(self, master=None):
        super().__init__(master, fg_color="transparent")
        self.cor_azul = "#0057B8"
        self.cor_verde = "#2EAF4A"
        self.arquivos_selecionados = []

        self.checkboxes_obs_cliente = []
        self.checkboxes_obs_adic = []

        self.mapeamento_obs = []
        self.selecionados_cliente = set()
        self.selecionados_adic = set()

        self.criar_interface()

    def criar_interface(self):
        # Cabeçalho
        header = ctk.CTkFrame(self, fg_color=self.cor_azul, corner_radius=0)
        header.pack(fill="x")

        try:
            logo = ctk.CTkImage(light_image=Image.open(CAMINHO_LOGO), size=(130, 30))
            lblLogo = ctk.CTkLabel(header, image=logo, text="")
            lblLogo.pack(side="left", padx=25, pady=12)
        except Exception:
            pass

        titulo = ctk.CTkLabel(
            header,
            text="CRM Formatação de Planilhas",
            text_color="white",
            font=("Segoe UI", 18, "bold")
        )
        titulo.pack(side="left", padx=15, pady=12)

        # Corpo Principal
        body = ctk.CTkFrame(self, fg_color="white", corner_radius=15)
        body.pack(padx=25, pady=20, fill="both", expand=True)

        # ==========================================
        # DIVISÃO EM DUAS COLUNAS
        # ==========================================
        container_colunas = ctk.CTkFrame(body, fg_color="transparent")
        container_colunas.pack(fill="both", expand=True, padx=20, pady=(15, 0))

        # --- Coluna da Esquerda (Arquivos e Datas) ---
        col_esq = ctk.CTkFrame(container_colunas, fg_color="transparent")
        col_esq.pack(side="left", fill="both", expand=True, padx=(0, 15))

        lbl1 = ctk.CTkLabel(col_esq, text=" 📂 Selecione os arquivos Excel (.xlsx)", font=("Segoe UI", 16, "bold"), text_color="#1C2D42")
        lbl1.pack(anchor="w", pady=(10, 5))

        self.btnArquivo = ctk.CTkButton(col_esq, text="Selecionar Arquivos (Lote)", height=45, corner_radius=8, font=("Segoe UI", 13, "bold"), command=self.abrir_arquivo)
        self.btnArquivo.pack(fill="x")

        self.lblArquivoStatus = ctk.CTkLabel(col_esq, text="Nenhum arquivo selecionado", text_color="gray")
        self.lblArquivoStatus.pack(anchor="w", pady=10)

        lbl2 = ctk.CTkLabel(col_esq, text=" 📅 Escolha a data de Programação", font=("Segoe UI", 16, "bold"), text_color="#1C2D42")
        lbl2.pack(anchor="w", pady=(80, 5))

        self.combo_datas = ctk.CTkComboBox(col_esq, values=[], height=40, state="disabled")
        self.combo_datas.pack(fill="x")

        # Dica visual para ocupar espaço na esquerda
        lbl_dica = ctk.CTkLabel(col_esq, text="Dica: O sistema lerá todas as datas e observações disponíveis nos arquivos\nselecionados automaticamente.", font=("Segoe UI", 12, "italic"), text_color="gray60", justify="left")
        lbl_dica.pack(anchor="w", pady=(30, 0))

        # --- Divisória Vertical ---
        divisoria = ctk.CTkFrame(container_colunas, width=2, fg_color="#E0E6ED")
        divisoria.pack(side="left", fill="y", padx=18, pady=10)

        # --- Coluna da Direita (Filtros de Observação) ---
        col_dir = ctk.CTkFrame(container_colunas, fg_color="transparent")
        col_dir.pack(side="right", fill="both", expand=True, padx=(15, 0))

        lbl3 = ctk.CTkLabel(col_dir, text=" ✔ Filtros Dinâmicos de Observação", font=("Segoe UI", 16, "bold"), text_color="#1C2D42")
        lbl3.pack(anchor="w", pady=(10, 5))

        frames_obs_container = ctk.CTkFrame(col_dir, fg_color="transparent")
        frames_obs_container.pack(fill="both", expand=True)

        # Lista de Cliente
        frame_cliente = ctk.CTkFrame(frames_obs_container, fg_color="transparent")
        frame_cliente.pack(side="left", fill="both", expand=True, padx=(0, 5))

        lbl_cli = ctk.CTkLabel(frame_cliente, text="Obs. Cliente", font=("Segoe UI", 13, "bold"), text_color="#333333")
        lbl_cli.pack(anchor="w")

        self.chk_all_cliente = ctk.CTkCheckBox(frame_cliente, text="Selecionar Todas", font=("Segoe UI", 12, "bold"), command=self.toggle_all_cliente)
        self.chk_all_cliente.pack(anchor="w", pady=(5, 5))

        # Como tiramos o 'height=100', ele vai expandir maravilhosamente pelo espaço restante
        self.scroll_obs_cliente = ctk.CTkScrollableFrame(frame_cliente, fg_color="#F0F2F5", corner_radius=8)
        self.scroll_obs_cliente.pack(fill="both", expand=True, pady=(0, 10))
        lbl_vazio_cli = ctk.CTkLabel(self.scroll_obs_cliente, text="Aguardando...", text_color="gray")
        lbl_vazio_cli.pack(pady=20)

        # Lista Adicional
        frame_adic = ctk.CTkFrame(frames_obs_container, fg_color="transparent")
        frame_adic.pack(side="left", fill="both", expand=True, padx=(5, 0))

        lbl_adi = ctk.CTkLabel(frame_adic, text="Obs. Adicionais", font=("Segoe UI", 13, "bold"), text_color="#333333")
        lbl_adi.pack(anchor="w")

        self.chk_all_adic = ctk.CTkCheckBox(frame_adic, text="Selecionar Todas", font=("Segoe UI", 12, "bold"), command=self.toggle_all_adic)
        self.chk_all_adic.pack(anchor="w", pady=(5, 5))

        self.scroll_obs_adic = ctk.CTkScrollableFrame(frame_adic, fg_color="#F0F2F5", corner_radius=8)
        self.scroll_obs_adic.pack(fill="both", expand=True, pady=(0, 10))
        lbl_vazio_adi = ctk.CTkLabel(self.scroll_obs_adic, text="Aguardando...", text_color="gray")
        lbl_vazio_adi.pack(pady=20)

        # ==========================================
        # RODAPÉ DE EXECUÇÃO (Spanning total na parte inferior)
        # ==========================================
        rodape_execucao = ctk.CTkFrame(body, fg_color="transparent")
        rodape_execucao.pack(fill="x", side="bottom", padx=25, pady=(25, 20), before=container_colunas)

        self.btnExecutar = ctk.CTkButton(
            rodape_execucao,
            width=260,
            height=55,
            corner_radius=10,
            fg_color="#BDC3C7",
            text_color="white",
            text="EXECUTAR FORMATAÇÃO EM LOTE",
            font=("Segoe UI", 16, "bold"),
            command=self.executar_processamento,
            state="disabled"
        )
        self.btnExecutar.pack(pady=(15, 10), fill="x")

        progress_frame = ctk.CTkFrame(rodape_execucao, fg_color="transparent")
        progress_frame.pack(fill="x", pady=(0, 5))

        self.progress = ctk.CTkProgressBar(progress_frame, height=14, progress_color=self.cor_verde, fg_color="#E0E6ED")
        self.progress.pack(side="left", fill="x", expand=True, padx=(0, 15))
        self.progress.set(0)

        self.lbl_porcentagem = ctk.CTkLabel(progress_frame, text="0%", font=("Segoe UI", 14, "bold"), text_color=self.cor_azul)
        self.lbl_porcentagem.pack(side="right")

        info_frame = ctk.CTkFrame(rodape_execucao, fg_color="transparent")
        info_frame.pack(fill="x")

        self.status = ctk.CTkLabel(info_frame, text="Status: Aguardando arquivos...", text_color="gray40")
        self.status.pack(side="left")

        footer = ctk.CTkLabel(info_frame, text="Desenvolvido para Operação GREE LOGÍSTICA", text_color="gray50", font=("Segoe UI", 10, "italic"))
        footer.pack(side="right")

    # ==========================
    # AS FUNÇÕES ABAIXO NÃO SOFRERAM NENHUMA ALTERAÇÃO
    # ==========================
    def on_date_change(self, nova_data):
        self.selecionados_cliente.clear()
        self.selecionados_adic.clear()
        for row in self.mapeamento_obs:
            if row['data'] == nova_data:
                if "nf" in row['cliente'].lower():
                    self.selecionados_cliente.add(row['cliente'])
                if "nf" in row['adic'].lower():
                    self.selecionados_adic.add(row['adic'])

        self.chk_all_cliente.deselect()
        self.chk_all_adic.deselect()
        self.atualizar_painel_filtros()

    def cb_cliente_clicked(self):
        self.selecionados_cliente = {cb.cget("text") for cb in self.checkboxes_obs_cliente if cb.get()}
        self.atualizar_painel_filtros()

    def cb_adic_clicked(self):
        self.selecionados_adic = {cb.cget("text") for cb in self.checkboxes_obs_adic if cb.get()}
        self.atualizar_painel_filtros()

    def toggle_all_cliente(self):
        estado = self.chk_all_cliente.get()
        for cb in self.checkboxes_obs_cliente:
            if estado:
                cb.select()
            else:
                cb.deselect()
        self.cb_cliente_clicked()

    def toggle_all_adic(self):
        estado = self.chk_all_adic.get()
        for cb in self.checkboxes_obs_adic:
            if estado:
                cb.select()
            else:
                cb.deselect()
        self.cb_adic_clicked()

    def atualizar_painel_filtros(self):
        data_escolhida = self.combo_datas.get()
        if not data_escolhida: return

        valid_clientes = set()
        valid_adics = set()

        for row in self.mapeamento_obs:
            if row['data'] == data_escolhida:
                if not self.selecionados_adic or row['adic'] in self.selecionados_adic:
                    valid_clientes.add(row['cliente'])

                if not self.selecionados_cliente or row['cliente'] in self.selecionados_cliente:
                    valid_adics.add(row['adic'])

        for widget in self.scroll_obs_cliente.winfo_children():
            widget.destroy()
        self.checkboxes_obs_cliente.clear()

        if valid_clientes:
            for obs in sorted(list(valid_clientes)):
                cb = ctk.CTkCheckBox(self.scroll_obs_cliente, text=obs, font=("Segoe UI", 12), command=self.cb_cliente_clicked)
                cb.pack(anchor="w", padx=10, pady=5)
                if obs in self.selecionados_cliente:
                    cb.select()
                self.checkboxes_obs_cliente.append(cb)
        else:
            lbl = ctk.CTkLabel(self.scroll_obs_cliente, text="Nenhum dado compatível.", text_color="gray")
            lbl.pack(pady=20)

        self.selecionados_cliente.intersection_update(valid_clientes)

        for widget in self.scroll_obs_adic.winfo_children():
            widget.destroy()
        self.checkboxes_obs_adic.clear()

        if valid_adics:
            for obs in sorted(list(valid_adics)):
                cb = ctk.CTkCheckBox(self.scroll_obs_adic, text=obs, font=("Segoe UI", 12), command=self.cb_adic_clicked)
                cb.pack(anchor="w", padx=10, pady=5)
                if obs in self.selecionados_adic:
                    cb.select()
                self.checkboxes_obs_adic.append(cb)
        else:
            lbl = ctk.CTkLabel(self.scroll_obs_adic, text="Nenhum dado compatível.", text_color="gray")
            lbl.pack(pady=20)

        self.selecionados_adic.intersection_update(valid_adics)

    def abrir_arquivo(self):
        caminhos = filedialog.askopenfilenames(title="Selecione as planilhas", filetypes=[("Arquivos Excel", "*.xlsx")])
        if caminhos:
            self.arquivos_selecionados = caminhos
            quantidade = len(caminhos)
            self.lblArquivoStatus.configure(text=f"{quantidade} arquivo(s) selecionado(s)", text_color="green")
            self.status.configure(text="Status: Analisando arquivos, datas e observações...")
            self.update_idletasks()
            self.carregar_lista_datas_e_obs()

    def carregar_lista_datas_e_obs(self):
        try:
            self.mapeamento_obs = []
            datas = set()

            for caminho in self.arquivos_selecionados:
                wb = openpyxl.load_workbook(caminho, data_only=True)
                ws = wb.active

                cabecalhos = [str(ws.cell(row=1, column=c).value).strip().lower() for c in range(1, ws.max_column + 1)]
                idx_cliente = -1
                idx_adic = -1

                for i, cab in enumerate(cabecalhos, 1):
                    cab_clean = cab.replace(".", "").replace(" ", "").strip()
                    if "cliente" in cab_clean and "obs" in cab_clean:
                        idx_cliente = i
                    elif "adicionais" in cab_clean and ("obs" in cab_clean or "observa" in cab_clean):
                        idx_adic = i

                if idx_cliente == -1 or idx_adic == -1:
                    for i, cab in enumerate(cabecalhos, 1):
                        cab_clean = cab.replace(".", "").replace(" ", "").strip()
                        if "obs" in cab_clean or "observa" in cab_clean:
                            if i != idx_cliente and i != idx_adic:
                                if idx_cliente == -1: idx_cliente = i
                                elif idx_adic == -1: idx_adic = i

                for i in range(2, ws.max_row + 1):
                    v = ws.cell(row=i, column=3).value
                    if v:
                        d = v.strftime('%d/%m/%Y') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()
                        if d:
                            datas.add(d)
                            val_cli = ws.cell(row=i, column=idx_cliente).value if idx_cliente != -1 else None
                            txt_cli = str(val_cli).strip() if val_cli is not None else ""
                            txt_cli_check = txt_cli if txt_cli else "(Vazio)"

                            val_adi = ws.cell(row=i, column=idx_adic).value if idx_adic != -1 else None
                            txt_adi = str(val_adi).strip() if val_adi is not None else ""
                            txt_adi_check = txt_adi if txt_adi else "(Vazio)"

                            self.mapeamento_obs.append({
                                'data': d,
                                'cliente': txt_cli_check,
                                'adic': txt_adi_check
                            })

                wb.close()

            if datas:
                lista_ordenada = sorted(list(datas))
                self.combo_datas.configure(values=lista_ordenada, state="readonly", command=self.on_date_change)
                self.combo_datas.set(lista_ordenada[0])

                self.btnExecutar.configure(state="normal", fg_color=self.cor_verde, hover_color="#22953D")

                self.status.configure(text="Status: Arquivos prontos para processamento.")
                self.on_date_change(lista_ordenada[0])
            else:
                messagebox.showwarning("Aviso", "Nenhuma data encontrada na Coluna C dos arquivos.")
                self.status.configure(text="Status: Aguardando datas válidas.")
                self.selecionados_cliente.clear()
                self.selecionados_adic.clear()
                self.atualizar_painel_filtros()

        except Exception as e:
            messagebox.showerror("Erro de Leitura", f"Falha ao processar arquivos: {e}")
            self.status.configure(text="Status: Erro na leitura dos arquivos.")

    def executar_processamento(self):
        data_escolhida = self.combo_datas.get()
        if not data_escolhida: return

        obs_cliente_selecionadas = [cb.cget("text") for cb in self.checkboxes_obs_cliente if cb.get()]
        obs_adic_selecionadas = [cb.cget("text") for cb in self.checkboxes_obs_adic if cb.get()]

        if not obs_cliente_selecionadas and not obs_adic_selecionadas and (self.checkboxes_obs_cliente or self.checkboxes_obs_adic):
            resp = messagebox.askyesno("Aviso", "Você não marcou nenhuma observação como filtro.\nDeseja processar TODAS as linhas da planilha ignorando essa separação?")
            if not resp: return

        self.btnExecutar.configure(state="disabled", fg_color="#BDC3C7")
        self.btnArquivo.configure(state="disabled")

        total_arquivos = len(self.arquivos_selecionados)
        sucessos = 0
        arquivo_atual = ""

        try:
            for idx, caminho in enumerate(self.arquivos_selecionados, 1):
                arquivo_atual = os.path.basename(caminho)
                self.status.configure(text=f"Status: Processando {idx}/{total_arquivos} ({arquivo_atual})")

                wb_teste = openpyxl.load_workbook(caminho, data_only=True)
                ws_teste = wb_teste.active
                cabecalhos = [str(ws_teste.cell(row=1, column=c).value).strip() for c in range(1, ws_teste.max_column + 1)]
                wb_teste.close()

                tipo_planilha = "PADRÃO"
                if "Código Unidade" in cabecalhos or "Unidade" in cabecalhos:
                    tipo_planilha = "G-MAX"

                sugestao_nome = arquivo_atual.replace(".xlsx", f"_{tipo_planilha}_FORMATADO.xlsx")
                caminho_final = ""

                while True:
                    caminho_final = filedialog.asksaveasfilename(
                        title=f"Salvar '{arquivo_atual}' como...",
                        initialfile=sugestao_nome,
                        defaultextension=".xlsx",
                        filetypes=[("Arquivos Excel", "*.xlsx")],
                        confirmoverwrite=False
                    )

                    if not caminho_final:
                        break

                    if os.path.exists(caminho_final):
                        messagebox.showwarning(
                            "Arquivo Já Existe",
                            f"O arquivo '{os.path.basename(caminho_final)}' já existe na pasta selecionada.\n"
                            "Por favor, escolha outro nome ou mude o local de gravação."
                        )
                    else:
                        break

                if not caminho_final:
                    continue

                self.status.configure(text=f"Status: Aplicando regras {tipo_planilha} em {arquivo_atual}")
                self.update_idletasks()

                if tipo_planilha == "G-MAX":
                    self.processar_gmax(caminho, data_escolhida, caminho_final, idx, total_arquivos, obs_cliente_selecionadas, obs_adic_selecionadas)
                else:
                    self.processar_padrao(caminho, data_escolhida, caminho_final, idx, total_arquivos, obs_cliente_selecionadas, obs_adic_selecionadas)
                sucessos += 1

            self.progress.set(1.0)
            self.lbl_porcentagem.configure(text="100%")
            self.status.configure(text="Status: Processamento concluído com sucesso!")
            messagebox.showinfo("Sucesso", f"Processamento concluído!\n{sucessos} arquivo(s) salvo(s).")

        except Exception as e:
            messagebox.showerror("Erro Crítico", f"Erro no arquivo: {arquivo_atual}\nErro: {str(e)}")
            self.status.configure(text="Status: Processamento interrompido com erro.")
        finally:
            self.btnExecutar.configure(state="normal", fg_color=self.cor_verde, hover_color="#22953D")
            self.btnArquivo.configure(state="normal")

    def processar_padrao(self, arquivo_path, data_escolhida, caminho_final, idx_arquivo, total_arquivos, obs_cliente_selecionadas, obs_adic_selecionadas):
        progresso_base = (idx_arquivo - 1) / total_arquivos
        progresso_por_arquivo = 1 / total_arquivos

        def atualizar_carga(porcentagem_interna):
            progresso_atual = progresso_base + (progresso_por_arquivo * porcentagem_interna)
            self.progress.set(progresso_atual)
            self.lbl_porcentagem.configure(text=f"{int(progresso_atual * 100)}%")
            self.update_idletasks()

        wb = openpyxl.load_workbook(arquivo_path)
        ws = wb.active
        qtd_colunas_originais = ws.max_column

        cabecalhos = [str(ws.cell(row=1, column=c).value).strip().lower() for c in range(1, ws.max_column + 1)]
        idx_obs_cliente = -1
        idx_obs_adc = -1

        for i, cab in enumerate(cabecalhos):
            cab_clean = cab.replace(".", "").replace(" ", "").strip()
            if "cliente" in cab_clean and "obs" in cab_clean:
                idx_obs_cliente = i
            elif "adicionais" in cab_clean and ("obs" in cab_clean or "observa" in cab_clean):
                idx_obs_adc = i

        if idx_obs_cliente == -1 or idx_obs_adc == -1:
            for i, cab in enumerate(cabecalhos):
                cab_clean = cab.replace(".", "").replace(" ", "").strip()
                if "obs" in cab_clean or "observa" in cab_clean:
                    if i != idx_obs_cliente and i != idx_obs_adc:
                        if idx_obs_cliente == -1: idx_obs_cliente = i
                        elif idx_obs_adc == -1: idx_obs_adc = i

        ws_copia = wb.copy_worksheet(ws)
        ws_copia.title = "Programação Completo"
        ws.title = "Programação Filtrada"

        linhas_raw = list(ws.iter_rows(min_row=2, values_only=True))
        linhas_validas = []

        req_len = max(ws.max_column, 65)
        ultimo_valor_conhecido = {i: None for i in range(9, 14)}

        atualizar_carga(0.2)

        for r_tuple in linhas_raw:
            r = list(r_tuple)
            if len(r) < req_len:
                r.extend([None] * (req_len - len(r)))

            val_nf = ""
            v_cliente = r[idx_obs_cliente] if idx_obs_cliente != -1 and idx_obs_cliente < len(r) else None
            v_adc = r[idx_obs_adc] if idx_obs_adc != -1 and idx_obs_adc < len(r) else None

            txt_cliente = str(v_cliente).strip() if v_cliente is not None else ""
            txt_adc = str(v_adc).strip() if v_adc is not None else ""

            txt_cliente_check = txt_cliente if txt_cliente else "(Vazio)"
            txt_adc_check = txt_adc if txt_adc else "(Vazio)"

            pular_linha = False

            if obs_cliente_selecionadas or obs_adic_selecionadas:
                match_cliente = True
                match_adic = True

                if obs_cliente_selecionadas:
                    match_cliente = (txt_cliente_check in obs_cliente_selecionadas)
                if obs_adic_selecionadas:
                    match_adic = (txt_adc_check in obs_adic_selecionadas)

                if not (match_cliente and match_adic):
                    pular_linha = True
                else:
                    if "nf" in txt_cliente.lower(): val_nf = txt_cliente_check
                    elif "nf" in txt_adc.lower(): val_nf = txt_adc_check
                    else:
                        val_nf = txt_cliente_check if txt_cliente_check != "(Vazio)" else txt_adc_check
                        if val_nf == "(Vazio)": val_nf = ""
            else:
                if "nf" in txt_cliente.lower(): val_nf = txt_cliente
                elif "nf" in txt_adc.lower(): val_nf = txt_adc
                else: val_nf = txt_cliente if txt_cliente else txt_adc

            if pular_linha:
                continue

            for i in range(9, 14):
                if r[i] is None or str(r[i]).strip() == "":
                    r[i] = ultimo_valor_conhecido[i]
                else:
                    ultimo_valor_conhecido[i] = r[i]

            r[0] = converter_numero(r[0])
            for i in [26, 32, 39, 42]:
                r[i] = moeda_americana_para_brasileira(r[i])
            r[30] = converter_cubagem(r[30])

            r[17] = r[50]
            r[18] = r[51]

            val_al = r[37]
            r[37] = remover_sufixo(val_al) if val_al and str(val_al).strip() != "" else r[33]
            r[40] = remover_sufixo(r[40])

            v = r[2]
            d = ""
            if v:
                d = v.strftime('%d/%m/%Y') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()

            if d == data_escolhida:
                try:
                    r[2] = datetime.datetime.strptime(data_escolhida, '%d/%m/%Y')
                except Exception:
                    pass
                linhas_validas.append((r, val_nf))

        atualizar_carga(0.4)

        def chave_ord_padrao(item):
            row, v_nf = item
            v_a = "" if row[0] is None else str(row[0]).strip()
            v_e = "" if row[7] is None else str(row[7]).strip()
            v_f = "" if row[13] is None else str(row[13]).strip()
            v_nf_str = "" if v_nf is None else str(v_nf).strip()
            return f"{v_a}|{v_e}|{v_f}|{v_nf_str}"

        linhas_validas.sort(key=chave_ord_padrao)

        atualizar_carga(0.5)

        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

        for item in linhas_validas:
            ws.append(item[0][:qtd_colunas_originais])

        for row_idx in range(2, len(linhas_validas) + 2):
            ws[f"C{row_idx}"].number_format = 'dd/mm/yyyy'
            for col_letter in ["AA", "AG", "AN", "AQ"]:
                ws[f"{col_letter}{row_idx}"].number_format = '#,##0.00'
            ws[f"AE{row_idx}"].number_format = '0.00'

        atualizar_carga(0.6)

        ws_drive = wb.create_sheet(title="DRIVE")
        colunas_padrao = [1, 3, 59, None, 8, 14, 15, 16, 18, 19, 20, 34, 36, 27, 28, 31, 38, 40, 41, 43, 10, 11, 12, 13, None, None, 54, 57, 60]

        borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        fonte_padrao = Font(name='Calibri', size=11)
        fonte_cabecalho = Font(name='Calibri', size=11, bold=True)
        alinhamento_central = Alignment(horizontal='center', vertical='center')

        max_lengths = [0] * len(colunas_padrao)

        for col_drive_idx, col_orig_idx in enumerate(colunas_padrao, start=1):
            celula_cabecalho = ws_drive.cell(row=1, column=col_drive_idx)
            if col_orig_idx is not None:
                val = ws.cell(row=1, column=col_orig_idx).value
                celula_cabecalho.value = val
                if val:
                    max_lengths[col_drive_idx - 1] = max(max_lengths[col_drive_idx - 1], len(str(val)))
            celula_cabecalho.font = fonte_cabecalho
            celula_cabecalho.border = borda_fina
            celula_cabecalho.alignment = alinhamento_central

        chave_anterior = None
        cor_index = -1
        cor_atual = "FFFF00"
        preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

        for row_idx, item in enumerate(linhas_validas, start=2):
            r, val_nf = item
            v_a = "" if r[0] is None else str(r[0]).strip()
            v_e = "" if r[7] is None else str(r[7]).strip()
            v_f = "" if r[13] is None else str(r[13]).strip()
            v_nf_str = "" if val_nf is None else str(val_nf).strip()

            chave = f"{v_a}|{v_e}|{v_f}|{v_nf_str}"

            if chave != chave_anterior:
                if chave_anterior is None:
                    cor_atual = "FFFF00"
                else:
                    cor_index = (cor_index + 1) % len(LISTA_CORES)
                    cor_atual = LISTA_CORES[cor_index]
                chave_anterior = chave
                preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

            nova_linha = []
            for col_idx_zero, col_orig_idx in enumerate(colunas_padrao):
                val = None
                if col_orig_idx is not None and (col_orig_idx - 1) < len(r):
                    val = r[col_orig_idx - 1]
                nova_linha.append(val)

                if val is not None:
                    tamanho = 10 if isinstance(val, (datetime.date, datetime.datetime)) else len(str(val))
                    if tamanho > max_lengths[col_idx_zero]:
                        max_lengths[col_idx_zero] = tamanho

            ws_drive.append(nova_linha)

            for col_idx, cel in enumerate(ws_drive[row_idx], start=1):
                cel.fill = preenchimento
                cel.font = fonte_padrao
                cel.border = borda_fina
                cel.alignment = alinhamento_central
                if col_idx == 2:
                    cel.number_format = 'dd/mm/yyyy'

            if row_idx % 50 == 0:
                atualizar_carga(0.6 + (0.3 * (row_idx / len(linhas_validas))))

        for idx, length in enumerate(max_lengths, start=1):
            if length > 0:
                ws_drive.column_dimensions[get_column_letter(idx)].width = length + 2

        atualizar_carga(0.9)
        wb.save(caminho_final)
        wb.close()

        self.status.configure(text="Status: Gerando Tabela Dinâmica...")
        self.update_idletasks()

        self.criar_dinamica_padrao_excel(caminho_final)
        atualizar_carga(1.0)

    def criar_dinamica_padrao_excel(self, arquivo_path):
        caminho_absoluto = os.path.abspath(arquivo_path)
        excel = None
        wb = None
        try:
            excel = abrir_excel_silencioso()
            if excel is None:
                messagebox.showerror("Erro de Dinâmica", "Não foi possível abrir o serviço do Excel em segundo plano (COM).\nVerifique se há planilhas travadas no seu Gerenciador de Tarefas.")
                return False

            excel, wb = abrir_workbook_com_retry(excel, caminho_absoluto)
            if wb is None:
                messagebox.showerror("Erro de Dinâmica", f"Não foi possível processar a planilha pelo Excel:\n{caminho_absoluto}")
                return False

            ws_dados = None
            for _ in range(5):
                try:
                    try:
                        ws_dados = wb.Worksheets("Programação Filtrada")
                    except Exception:
                        ws_dados = wb.Worksheets(1)
                    break
                except pywintypes.com_error as e:
                    if _codigo_com_error(e) == -2147418111:
                        time.sleep(1)
                    else:
                        raise

            for _ in range(5):
                try:
                    for sheet in list(wb.Worksheets):
                        if sheet.Name == "Dinâmicas Padrão":
                            sheet.Delete()
                    break
                except pywintypes.com_error as e:
                    if _codigo_com_error(e) == -2147418111:
                        time.sleep(1)
                    else:
                        raise

            ws_dinamica = wb.Worksheets.Add(After=ws_dados)
            ws_dinamica.Name = _normalizar_titulo_planilha("Dinâmicas Padrão")

            ultima_linha = ws_dados.Cells(ws_dados.Rows.Count, 1).End(-4162).Row
            ultima_coluna = ws_dados.Cells(1, ws_dados.Columns.Count).End(-4159).Column

            r1c1_address = f"'{ws_dados.Name}'!R1C1:R{ultima_linha}C{ultima_coluna}"

            pc = wb.PivotCaches().Create(SourceType=1, SourceData=r1c1_address)

            range_destino_1 = ws_dinamica.Range("A4")
            pt1 = pc.CreatePivotTable(TableDestination=range_destino_1, TableName="Dinamica1")
            pt1.TableStyle2 = "PivotStyleMedium6"
            pt1.RowAxisLayout(1)

            campo_ciclo = obter_campo_dinamico(pt1, 'Ciclo')
            if campo_ciclo:
                try: campo_ciclo.Orientation = 3
                except: pass

            campos_1 = ['Data de Programação', 'Loading', 'Frete', 'Razão Social', 'Cidade', 'UF', 'Transportador (Nome)']
            for i, nome_campo in enumerate(campos_1):
                f = obter_campo_dinamico(pt1, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            for nome, caption in [('Qty', 'Soma de Qty'), ('Cubagem M³', 'Soma de Cubagem M³'), ('Ttl Faturado R$', 'Soma de Ttl Faturado R$')]:
                f = obter_campo_dinamico(pt1, nome)
                if f:
                    try:
                        val = pt1.AddDataField(f)
                        val.Function, val.Caption, val.NumberFormat = -4157, caption, "General"
                    except: pass

            pt1.ColumnGrand = True
            pt1.RowGrand = False

            range_destino_2 = ws_dinamica.Range("M4")
            pt2 = pc.CreatePivotTable(TableDestination=range_destino_2, TableName="Dinamica2")
            pt2.TableStyle2 = "PivotStyleMedium6"
            pt2.RowAxisLayout(1)

            for i, nome_campo in enumerate(['Loading', 'Pedido', 'Razão Social', 'Código Evap.', 'Cód. Cond.']):
                f = obter_campo_dinamico(pt2, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            f_qty = obter_campo_dinamico(pt2, 'Qty')
            if f_qty:
                try:
                    val3 = pt2.AddDataField(f_qty)
                    val3.Function, val3.Caption, val3.NumberFormat = -4157, 'Total ', "General"
                except: pass

            pt2.ColumnGrand = True
            pt2.RowGrand = False
            ws_dinamica.Columns("A:P").AutoFit()

            wb.Save()
            return True
        except Exception as e:
            mensagem = f"Falha interna ao desenhar a Dinâmica Padrão.\n\nDetalhes:\n{str(e)}\n\nTraceback:\n{traceback.format_exc()}"
            messagebox.showerror("Erro Dinâmicas Padrão", mensagem)
            return False
        finally:
            if wb:
                try: wb.Close(SaveChanges=False)
                except: pass
            if excel:
                try: excel.Quit()
                except: pass
            try: pythoncom.CoUninitialize()
            except Exception: pass

    def processar_gmax(self, arquivo_path, data_escolhida, caminho_final, idx_arquivo, total_arquivos, obs_cliente_selecionadas, obs_adic_selecionadas):
        progresso_base = (idx_arquivo - 1) / total_arquivos
        progresso_por_arquivo = 1 / total_arquivos

        def atualizar_carga(porcentagem_interna):
            progresso_atual = progresso_base + (progresso_por_arquivo * porcentagem_interna)
            self.progress.set(progresso_atual)
            self.lbl_porcentagem.configure(text=f"{int(progresso_atual * 100)}%")
            self.update_idletasks()

        wb = openpyxl.load_workbook(arquivo_path)
        ws = wb.active
        qtd_colunas_originais = ws.max_column

        cabecalhos = [str(ws.cell(row=1, column=c).value).strip().lower() for c in range(1, ws.max_column + 1)]
        idx_obs_cliente = -1
        idx_obs_adc = -1

        for i, cab in enumerate(cabecalhos):
            cab_clean = cab.replace(".", "").replace(" ", "").strip()
            if "cliente" in cab_clean and "obs" in cab_clean:
                idx_obs_cliente = i
            elif "adicionais" in cab_clean and ("obs" in cab_clean or "observa" in cab_clean):
                idx_obs_adc = i

        if idx_obs_cliente == -1 or idx_obs_adc == -1:
            for i, cab in enumerate(cabecalhos):
                cab_clean = cab.replace(".", "").replace(" ", "").strip()
                if "obs" in cab_clean or "observa" in cab_clean:
                    if i != idx_obs_cliente and i != idx_obs_adc:
                        if idx_obs_cliente == -1: idx_obs_cliente = i
                        elif idx_obs_adc == -1: idx_obs_adc = i

        ws_copia = wb.copy_worksheet(ws)
        ws_copia.title = "Programação Completo"
        ws.title = "Programação Filtrada"

        linhas_raw = list(ws.iter_rows(min_row=2, values_only=True))
        linhas_validas = []
        req_len = max(ws.max_column, 50)

        atualizar_carga(0.2)

        for r_tuple in linhas_raw:
            r = list(r_tuple)
            if len(r) < req_len:
                r.extend([None] * (req_len - len(r)))

            val_nf = ""
            v_cliente = r[idx_obs_cliente] if idx_obs_cliente != -1 and idx_obs_cliente < len(r) else None
            v_adc = r[idx_obs_adc] if idx_obs_adc != -1 and idx_obs_adc < len(r) else None

            txt_cliente = str(v_cliente).strip() if v_cliente is not None else ""
            txt_adc = str(v_adc).strip() if v_adc is not None else ""

            txt_cliente_check = txt_cliente if txt_cliente else "(Vazio)"
            txt_adc_check = txt_adc if txt_adc else "(Vazio)"

            pular_linha = False

            if obs_cliente_selecionadas or obs_adic_selecionadas:
                match_cliente = True
                match_adic = True

                if obs_cliente_selecionadas:
                    match_cliente = (txt_cliente_check in obs_cliente_selecionadas)
                if obs_adic_selecionadas:
                    match_adic = (txt_adc_check in obs_adic_selecionadas)

                if not (match_cliente and match_adic):
                    pular_linha = True
                else:
                    if "nf" in txt_cliente.lower(): val_nf = txt_cliente_check
                    elif "nf" in txt_adc.lower(): val_nf = txt_adc_check
                    else:
                        val_nf = txt_cliente_check if txt_cliente_check != "(Vazio)" else txt_adc_check
                        if val_nf == "(Vazio)": val_nf = ""
            else:
                if "nf" in txt_cliente.lower(): val_nf = txt_cliente
                elif "nf" in txt_adc.lower(): val_nf = txt_adc
                else: val_nf = txt_cliente if txt_cliente else txt_adc

            if pular_linha:
                continue

            v_ad = r[29]
            if v_ad is not None:
                try: r[29] = float(str(v_ad).strip().replace(',', '.'))
                except ValueError: pass

            for i in [30, 31]:
                v = r[i]
                if v is not None:
                    texto = str(v).strip()
                    if ',' in texto and '.' in texto: texto = texto.replace(',', '')
                    elif ',' in texto: texto = texto.replace(',', '.')
                    try: r[i] = float(texto)
                    except ValueError: pass

            r[17] = r[38]
            r[18] = r[39]
            r[34] = remover_sufixo(r[34])

            v = r[2]
            d = ""
            if v:
                d = v.strftime('%d/%m/%Y') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()

            if d == data_escolhida:
                try:
                    r[2] = datetime.datetime.strptime(data_escolhida, '%d/%m/%Y')
                except Exception:
                    pass
                linhas_validas.append((r, val_nf))

        atualizar_carga(0.4)

        def chave_ord_gmax(item):
            row, v_nf = item
            v_a = "" if row[0] is None else str(row[0]).strip()
            v_e = "" if row[7] is None else str(row[7]).strip()
            v_f = "" if row[13] is None else str(row[13]).strip()
            v_ac = "" if row[46] is None else str(row[46]).strip()
            v_nf_str = "" if v_nf is None else str(v_nf).strip()
            return f"{v_a}|{v_e}|{v_f}|{v_ac}|{v_nf_str}"

        linhas_validas.sort(key=chave_ord_gmax)

        atualizar_carga(0.5)

        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

        for item in linhas_validas:
            ws.append(item[0][:qtd_colunas_originais])

        for row_idx in range(2, len(linhas_validas) + 2):
            ws[f"C{row_idx}"].number_format = 'dd/mm/yyyy'
            ws[f"AD{row_idx}"].number_format = '0.00'
            for col_letter in ["AE", "AF"]:
                ws[f"{col_letter}{row_idx}"].number_format = '#,##0.00'

        atualizar_carga(0.6)

        ws_drive = wb.create_sheet(title="DRIVE")
        colunas_gmax = [1, 3, 46, None, 8, 14, 15, 16, 18, 19, 22, 33, 37, 31, 27, 30, 35, 31, 34, None, 10, 11, 12, 13, None, None, 41, 44, 47]

        borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        fonte_padrao = Font(name='Calibri', size=11)
        fonte_cabecalho = Font(name='Calibri', size=11, bold=True)
        alinhamento_central = Alignment(horizontal='center', vertical='center')

        max_lengths = [0] * len(colunas_gmax)

        for col_drive_idx, col_orig_idx in enumerate(colunas_gmax, start=1):
            celula_cabecalho = ws_drive.cell(row=1, column=col_drive_idx)
            if col_orig_idx is not None:
                val = ws.cell(row=1, column=col_orig_idx).value
                celula_cabecalho.value = val
                if val:
                    max_lengths[col_drive_idx - 1] = max(max_lengths[col_drive_idx - 1], len(str(val)))
            celula_cabecalho.font = fonte_cabecalho
            celula_cabecalho.border = borda_fina
            celula_cabecalho.alignment = alinhamento_central

        chave_anterior = None
        cor_index = -1
        cor_atual = "FFFF00"
        preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

        for row_idx, item in enumerate(linhas_validas, start=2):
            r, val_nf = item
            v_a = "" if r[0] is None else str(r[0]).strip()
            v_e = "" if r[7] is None else str(r[7]).strip()
            v_f = "" if r[13] is None else str(r[13]).strip()
            v_ac = "" if r[46] is None else str(r[46]).strip()
            v_nf_str = "" if val_nf is None else str(val_nf).strip()

            chave = f"{v_a}|{v_e}|{v_f}|{v_ac}|{v_nf_str}"

            if chave != chave_anterior:
                if chave_anterior is None:
                    cor_atual = "FFFF00"
                else:
                    cor_index = (cor_index + 1) % len(LISTA_CORES)
                    cor_atual = LISTA_CORES[cor_index]
                chave_anterior = chave
                preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

            nova_linha = []
            for col_idx_zero, col_orig_idx in enumerate(colunas_gmax):
                val = None
                if col_orig_idx is not None and (col_orig_idx - 1) < len(r):
                    val = r[col_orig_idx - 1]
                nova_linha.append(val)

                if val is not None:
                    tamanho = 10 if isinstance(val, (datetime.date, datetime.datetime)) else len(str(val))
                    if tamanho > max_lengths[col_idx_zero]:
                        max_lengths[col_idx_zero] = tamanho

            ws_drive.append(nova_linha)

            for col_idx, cel in enumerate(ws_drive[row_idx], start=1):
                cel.fill = preenchimento
                cel.font = fonte_padrao
                cel.border = borda_fina
                cel.alignment = alinhamento_central
                if col_idx == 2:
                    cel.number_format = 'dd/mm/yyyy'

            if row_idx % 50 == 0:
                atualizar_carga(0.6 + (0.3 * (row_idx / len(linhas_validas))))

        for idx, length in enumerate(max_lengths, start=1):
            if length > 0:
                ws_drive.column_dimensions[get_column_letter(idx)].width = length + 2

        atualizar_carga(0.9)
        wb.save(caminho_final)
        wb.close()

        self.status.configure(text="Status: Gerando Tabela Dinâmica...")
        self.update_idletasks()

        self.criar_dinamica_gmax_excel(caminho_final)
        atualizar_carga(1.0)

    def criar_dinamica_gmax_excel(self, arquivo_path):
        caminho_absoluto = os.path.abspath(arquivo_path)
        excel = None
        wb = None
        try:
            excel = abrir_excel_silencioso()
            if excel is None:
                messagebox.showerror("Erro de Dinâmica", "Não foi possível abrir o serviço do Excel em segundo plano (COM).\nVerifique se há planilhas travadas no seu Gerenciador de Tarefas.")
                return False

            excel, wb = abrir_workbook_com_retry(excel, caminho_absoluto)
            if wb is None:
                messagebox.showerror("Erro de Dinâmica", f"Não foi possível processar a planilha pelo Excel:\n{caminho_absoluto}")
                return False

            ws_dados = None
            for _ in range(5):
                try:
                    try:
                        ws_dados = wb.Worksheets("Programação Filtrada")
                    except Exception:
                        ws_dados = wb.Worksheets(1)
                    break
                except pywintypes.com_error as e:
                    if _codigo_com_error(e) == -2147418111:
                        time.sleep(1)
                    else:
                        raise

            for _ in range(5):
                try:
                    for sheet in list(wb.Worksheets):
                        if sheet.Name == "Dinâmica G-MAX":
                            sheet.Delete()
                    break
                except pywintypes.com_error as e:
                    if _codigo_com_error(e) == -2147418111:
                        time.sleep(1)
                    else:
                        raise

            ws_dinamica = wb.Worksheets.Add(After=ws_dados)
            ws_dinamica.Name = _normalizar_titulo_planilha("Dinâmica G-MAX")

            ultima_linha = ws_dados.Cells(ws_dados.Rows.Count, 1).End(-4162).Row
            ultima_coluna = ws_dados.Cells(1, ws_dados.Columns.Count).End(-4159).Column

            r1c1_address = f"'{ws_dados.Name}'!R1C1:R{ultima_linha}C{ultima_coluna}"

            pc = wb.PivotCaches().Create(SourceType=1, SourceData=r1c1_address)

            range_destino_1 = ws_dinamica.Range("A4")
            pt1 = pc.CreatePivotTable(TableDestination=range_destino_1, TableName="DinamicaGMAX1")
            pt1.TableStyle2 = "PivotStyleMedium6"
            pt1.RowAxisLayout(1)

            f_unidade = obter_campo_dinamico(pt1, 'Unidade')
            if f_unidade:
                try: f_unidade.Orientation = 3
                except: pass

            campos_gmax_1 = ['Data de Programação', 'Loading', 'Frete', 'Razão Social', 'Cidade', 'UF', 'Transportador']
            for i, nome_campo in enumerate(campos_gmax_1):
                f = obter_campo_dinamico(pt1, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            for nome, caption in [('Qty', 'Soma de Qty'), ('Cubagem M³', 'Soma de Cubagem M³'), ('Ttl. Faturado (R$)', 'Soma de Ttl. Faturado (R$)')]:
                f = obter_campo_dinamico(pt1, nome)
                if f:
                    try:
                        val = pt1.AddDataField(f)
                        val.Function, val.Caption, val.NumberFormat = -4157, caption, "General"
                    except: pass

            pt1.ColumnGrand = True
            pt1.RowGrand = False

            range_destino_2 = ws_dinamica.Range("M4")
            pt2 = pc.CreatePivotTable(TableDestination=range_destino_2, TableName="DinamicaGMAX2")
            pt2.TableStyle2 = "PivotStyleMedium6"
            pt2.RowAxisLayout(1)

            campos_gmax_2 = ['Loading', 'Pedido', 'Razão Social', 'Código Unidade', 'Unidade']
            for i, nome_campo in enumerate(campos_gmax_2):
                f = obter_campo_dinamico(pt2, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            f_qty = obter_campo_dinamico(pt2, 'Qty')
            if f_qty:
                try:
                    val3 = pt2.AddDataField(f_qty)
                    val3.Function, val3.Caption, val3.NumberFormat = -4157, 'Soma Qty ', "General"
                except: pass

            pt2.ColumnGrand = True
            pt2.RowGrand = False
            ws_dinamica.Columns("A:R").AutoFit()

            wb.Save()
            return True
        except Exception as e:
            mensagem = f"Falha interna ao desenhar a Dinâmica G-MAX.\n\nDetalhes:\n{str(e)}\n\nTraceback:\n{traceback.format_exc()}"
            messagebox.showerror("Erro Dinâmicas G-MAX", mensagem)
            return False
        finally:
            if wb:
                try: wb.Close(SaveChanges=False)
                except: pass
            if excel:
                try: excel.Quit()
                except: pass
            try: pythoncom.CoUninitialize()
            except Exception: pass

# ==========================================
# CÓDIGO 2: GERADOR DE COTAÇÕES (Frame Interno)
# ==========================================
class CotacaoApp(ctk.CTkFrame):
    def __init__(self, master=None):
        super().__init__(master, fg_color="transparent")
        self.azul_gree = "#0057B8"
        self.laranja_gree = "#F26522"
        self.verde_gree = "#2EAF4A"
        self.df_dados = None
        self.df_original_cif = None
        self.is_gmax = False
        self.setup_ui()

    def setup_ui(self):
        header = ctk.CTkFrame(self, fg_color=self.azul_gree, corner_radius=0)
        header.pack(fill="x", side="top")

        try:
            logo = ctk.CTkImage(light_image=Image.open(CAMINHO_LOGO), size=(130, 30))
            lblLogo = ctk.CTkLabel(header, image=logo, text="")
            lblLogo.pack(side="left", padx=25, pady=12)
        except Exception:
            pass

        title = ctk.CTkLabel(
            header,
            text="Gerador de Cotações",
            text_color="white",
            font=("Segoe UI", 18, "bold")
        )
        title.pack(side="left", padx=15, pady=12)

        body = ctk.CTkFrame(self, fg_color="white", corner_radius=15)
        body.pack(fill="both", expand=True, padx=25, pady=20)

        lbl1 = ctk.CTkLabel(
            body,
            text=" 📂 Importe a planilha de dados",
            font=("Segoe UI", 17, "bold"),
            text_color="#1C2D42"
        )
        lbl1.pack(anchor="w", padx=30, pady=(20, 8))

        self.btn_carregar = ctk.CTkButton(
            body,
            width=220,
            height=40,
            corner_radius=10,
            fg_color=self.laranja_gree,
            hover_color="#D35400",
            text="Importar Base de Dados",
            font=("Segoe UI", 13, "bold"),
            command=self.carregar_arquivo
        )
        self.btn_carregar.pack(padx=30, fill="x")

        self.lbl_status = tk.Label(
            body,
            text="Nenhum arquivo carregado",
            bg="white",
            fg="#7F8C8D",
            font=("Segoe UI", 12, "bold")
        )
        self.lbl_status.pack(anchor="w", padx=35, pady=5)

        lbl2 = ctk.CTkLabel(
            body,
            text=" 📅 Organize a ordem das Cotações",
            font=("Segoe UI", 17, "bold"),
            text_color="#1C2D42"
        )
        lbl2.pack(anchor="w", padx=30, pady=(15, 8))

        lbl_instrucao = ctk.CTkLabel(
            body,
            text="Arraste os itens na lista ou utilize os botões de subir e descer ao lado para ordenar:",
            font=("Segoe UI", 12, "italic"),
            text_color="gray50"
        )
        lbl_instrucao.pack(anchor="w", padx=35, pady=(0, 5))

        frame_meio = ctk.CTkFrame(body, fg_color="transparent")
        frame_meio.pack(fill="both", expand=True, padx=30, pady=5)

        frame_lista = ctk.CTkFrame(frame_meio, fg_color="white", border_color="#BDC3C7", border_width=1, corner_radius=8)
        frame_lista.pack(side="left", fill="both", expand=True)

        self.listbox = tk.Listbox(
            frame_lista,
            font=("Segoe UI", 12),
            selectbackground=self.azul_gree,
            selectforeground="white",
            activestyle="none",
            relief="flat",
            borderwidth=0,
            highlightthickness=0,
            bg="white",
            fg="#2C3E50"
        )
        self.listbox.pack(side="left", fill="both", expand=True, padx=8, pady=8)

        scrollbar = ttk.Scrollbar(frame_lista, orient="vertical", command=self.listbox.yview)
        scrollbar.pack(side="right", fill="y", padx=(0, 4), pady=4)
        self.listbox.config(yscrollcommand=scrollbar.set)

        frame_botoes = ctk.CTkFrame(frame_meio, fg_color="transparent")
        frame_botoes.pack(side="right", fill="y", padx=(15, 0))

        btn_up = ctk.CTkButton(
            frame_botoes,
            text="▲ Subir",
            command=self.mover_cima,
            fg_color="#EAF2FC",
            text_color=self.azul_gree,
            hover_color="#D0E1F9",
            font=("Segoe UI", 12, "bold"),
            width=110,
            height=40,
            corner_radius=8
        )
        btn_up.pack(pady=(20, 10))

        btn_down = ctk.CTkButton(
            frame_botoes,
            text="▼ Descer",
            command=self.mover_baixo,
            fg_color="#EAF2FC",
            text_color=self.azul_gree,
            hover_color="#D0E1F9",
            font=("Segoe UI", 12, "bold"),
            width=110,
            height=40,
            corner_radius=8
        )
        btn_down.pack(pady=10)

        self.listbox.bind("<Button-1>", self.on_drag_start)
        self.listbox.bind("<B1-Motion>", self.on_drag_motion)

        lbl3 = ctk.CTkLabel(
            body,
            text=" 📥 Gere a cotação final formatada",
            font=("Segoe UI", 17, "bold"),
            text_color="#1C2D42"
        )
        lbl3.pack(anchor="w", padx=30, pady=(15, 8))

        self.btn_gerar = ctk.CTkButton(
            body,
            width=260,
            height=55,
            corner_radius=10,
            fg_color="#BDC3C7",
            text_color="white",
            text="EXPORTAR COTAÇÃO FORMATADA",
            font=("Segoe UI", 16, "bold"),
            command=self.gerar_planilha,
            state="disabled"
        )
        self.btn_gerar.pack(padx=30, fill="x", pady=(0, 20))

        self.btn_gerar.ativar = lambda *args: self.btn_gerar.configure(
            state="normal",
            fg_color=self.verde_gree,
            hover_color="#22953D"
        )
        self.btn_gerar.desativar = lambda: self.btn_gerar.configure(
            state="disabled",
            fg_color="#BDC3C7"
        )

        footer = ctk.CTkLabel(
            self,
            text="Desenvolvido para Operação GREE LOGÍSTICA",
            font=("Segoe UI", 10, "italic"),
            text_color="gray50"
        )
        footer.pack(pady=8)

    def mover_cima(self):
        try:
            sel = self.listbox.curselection()[0]
            if sel > 0:
                text = self.listbox.get(sel)
                self.listbox.delete(sel)
                self.listbox.insert(sel - 1, text)
                self.listbox.select_set(sel - 1)
        except IndexError: pass

    def mover_baixo(self):
        try:
            sel = self.listbox.curselection()[0]
            if sel < self.listbox.size() - 1:
                text = self.listbox.get(sel)
                self.listbox.delete(sel)
                self.listbox.insert(sel + 1, text)
                self.listbox.select_set(sel + 1)
        except IndexError: pass

    def on_drag_start(self, event):
        self._drag_start_index = self.listbox.nearest(event.y)

    def on_drag_motion(self, event):
        i = self.listbox.nearest(event.y)
        if i < self.listbox.size() and i >= 0:
            if i != self._drag_start_index:
                item = self.listbox.get(self._drag_start_index)
                self.listbox.delete(self._drag_start_index)
                self.listbox.insert(i, item)
                self.listbox.select_clear(0, tk.END)
                self.listbox.select_set(i)
                self._drag_start_index = i

    def carregar_arquivo(self):
        caminho_entrada = filedialog.askopenfilename(
            title="Selecione a base de dados (Excel)",
            filetypes=[("Arquivos Excel", "*.xlsx *.xls")]
        )
        if not caminho_entrada: return

        try:
            xls = pd.ExcelFile(caminho_entrada)
            if "Programação Filtrada" not in xls.sheet_names:
                messagebox.showerror("Erro", "A aba 'Programação Filtrada' não foi encontrada.")
                return

            df_bruto = pd.read_excel(caminho_entrada, sheet_name="Programação Filtrada")

            col_frete = df_bruto.columns[letter_to_index('G')]
            df_cif = df_bruto[df_bruto[col_frete].astype(str).str.upper().str.strip() == 'CIF'].copy()

            if df_cif.empty:
                messagebox.showwarning("Aviso", "Nenhuma carga 'CIF' encontrada na base.")
                return

            self.df_original_cif = df_cif.copy()

            self.is_gmax = False
            if len(df_bruto.columns) > 32:
                col_nome_ag = str(df_bruto.columns[letter_to_index('AG')]).strip().lower()
                if 'venda' in col_nome_ag or 'cod' in col_nome_ag:
                    self.is_gmax = True

            visao_texto = "G-MAX" if self.is_gmax else "PADRÃO"

            col_tipo_veiculo = df_cif.columns[letter_to_index('F')]
            col_item = df_cif.columns[letter_to_index('A')]
            df_cif[col_tipo_veiculo] = df_cif.groupby(col_item)[col_tipo_veiculo].transform(lambda x: x.ffill().bfill())

            col_y_idx = letter_to_index('Y')
            df_cif['CICLO_calc'] = df_cif.iloc[:, col_y_idx].fillna('') if col_y_idx < len(df_bruto.columns) else ''

            if self.is_gmax:
                mapa = {'ITEM': 'A', 'CIDADE': 'R', 'UF': 'S', 'CLIENTE': 'P', 'CODIGO': 'AG', 'QTD': 'AA', 'M3': 'AD', 'VALOR': 'AF', 'DESCRICAO': 'AK'}
            else:
                mapa = {'ITEM': 'A', 'CIDADE': 'R', 'UF': 'S', 'CLIENTE': 'P', 'CODIGO': 'AH', 'QTD': 'AB', 'M3': 'AE', 'VALOR': 'AG', 'DESCRICAO': 'AJ'}

            df_cif['ITEM_calc'] = df_cif.iloc[:, letter_to_index(mapa['ITEM'])].apply(safe_str)
            df_cif['CIDADE_calc'] = df_cif.iloc[:, letter_to_index(mapa['CIDADE'])].fillna('')
            df_cif['UF_calc'] = df_cif.iloc[:, letter_to_index(mapa['UF'])].fillna('')
            df_cif['CLIENTE_calc'] = df_cif.iloc[:, letter_to_index(mapa['CLIENTE'])].fillna('')
            df_cif['CODIGO_calc'] = df_cif.iloc[:, letter_to_index(mapa['CODIGO'])].fillna('')
            df_cif['QTD_calc'] = df_cif.iloc[:, letter_to_index(mapa['QTD'])].fillna(0)
            df_cif['M3_calc'] = df_cif.iloc[:, letter_to_index(mapa['M3'])].fillna(0.0)
            df_cif['VALOR_calc'] = df_cif.iloc[:, letter_to_index(mapa['VALOR'])].fillna(0.0)
            df_cif['DESCRICAO_calc'] = df_cif.iloc[:, letter_to_index(mapa['DESCRICAO'])].fillna('')
            df_cif['TIPO_VEICULO_calc'] = df_cif[col_tipo_veiculo].fillna('')

            self.df_dados = df_cif

            self.lbl_status.config(text=f"Sucesso! Visão Identificada: {visao_texto} | Cargas CIF: {len(df_cif)}", fg="#27AE60")

            self.listbox.delete(0, tk.END)
            itens_unicos = self.df_dados['ITEM_calc'].drop_duplicates().tolist()
            for item in itens_unicos:
                if str(item).strip() != "":
                    self.listbox.insert(tk.END, str(item))

            self.btn_gerar.ativar(self.azul_gree)

        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao ler o arquivo:\n\n{str(e)}")

    def gerar_aba_excel(self, ws, df_filtrado, tipo_aba):
        fonte_branca_normal = Font(bold=False, color="FFFFFF")
        fonte_preta_normal = Font(bold=False, color="000000")
        fonte_preta_negrito = Font(bold=True, color="000000")

        fill_azul_escuro = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        fill_cinza = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        fill_amarelo = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        alinhamento_centro = Alignment(horizontal="center", vertical="center")

        borda_fina = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        if tipo_aba == "CABOTAGEM":
            transp = ["COSTA BRASIL", "CENTER CARGO", "TECMAR"]
        else:
            transp = ["TRANSBUIATTE", "SPEED", "SR LOG", "BUSSOLA", "TODOBRASIL", "AMAZON", "GAB", "BERTOLINI"]

        titulos = ["ITEM"] + transp + ["CIDADE", "UF", "CLIENTE", "CODIGO", "QTD", "M³", "Valor", "DESCRIÇÃO"]

        titulos_colunas = {i+1: nome for i, nome in enumerate(titulos)}

        num_transp = len(transp)

        col_cidade = titulos.index("CIDADE") + 1
        col_uf = titulos.index("UF") + 1
        col_cliente = titulos.index("CLIENTE") + 1
        col_codigo = titulos.index("CODIGO") + 1
        col_qtd = titulos.index("QTD") + 1
        col_m3 = titulos.index("M³") + 1
        col_valor = titulos.index("Valor") + 1
        col_descricao = titulos.index("DESCRIÇÃO") + 1

        linha_atual = 1
        ordem_desejada = list(self.listbox.get(0, tk.END))

        for item_nome in ordem_desejada:
            grupo = df_filtrado[df_filtrado['ITEM_calc'] == item_nome]
            if grupo.empty: continue

            for col_idx, col_nome in titulos_colunas.items():
                celula = ws.cell(row=linha_atual, column=col_idx, value=col_nome)
                celula.fill = fill_azul_escuro
                celula.font = fonte_branca_normal
                celula.alignment = alinhamento_centro
                celula.border = borda_fina

            linha_atual += 1
            linha_inicial = linha_atual

            for _, row in grupo.iterrows():
                ws.cell(row=linha_atual, column=col_cidade, value=row.get('CIDADE_calc', ''))
                ws.cell(row=linha_atual, column=col_uf, value=row.get('UF_calc', ''))
                ws.cell(row=linha_atual, column=col_cliente, value=row.get('CLIENTE_calc', ''))
                ws.cell(row=linha_atual, column=col_codigo, value=row.get('CODIGO_calc', ''))

                cel_qtd = ws.cell(row=linha_atual, column=col_qtd, value=row.get('QTD_calc', 0))
                cel_qtd.number_format = '0'

                cel_m3 = ws.cell(row=linha_atual, column=col_m3, value=row.get('M3_calc', 0.0))
                cel_m3.number_format = '0.00'

                is_painel = False
                if not self.is_gmax:
                    ciclo_info = str(row.get('CICLO_calc', '')).strip().upper()
                    if 'PAINEL' in ciclo_info:
                        is_painel = True

                if is_painel:
                    cel_valor_cel = ws.cell(row=linha_atual, column=col_valor, value="")
                else:
                    cel_valor_cel = ws.cell(row=linha_atual, column=col_valor, value=row.get('VALOR_calc', 0.0))
                    cel_valor_cel.number_format = '"R$" #,##0.00'

                ws.cell(row=linha_atual, column=col_descricao, value=row.get('DESCRICAO_calc', ''))

                for c in range(1, len(titulos) + 1):
                    cel = ws.cell(row=linha_atual, column=c)
                    cel.border = borda_fina
                    if c >= col_cidade: cel.alignment = alinhamento_centro

                linha_atual += 1

            linha_final = linha_atual - 1
            linha_total = linha_atual

            ws.cell(row=linha_inicial, column=1, value=item_nome)
            letra_valor = get_column_letter(col_valor)

            for col in range(1, num_transp + 2):
                ws.merge_cells(start_row=linha_inicial, start_column=col, end_row=linha_final, end_column=col)
                cel_mesclada = ws.cell(row=linha_inicial, column=col)

                for r in range(linha_inicial, linha_final + 1):
                    c_borda = ws.cell(row=r, column=col)
                    c_borda.border = borda_fina
                    c_borda.alignment = alinhamento_centro

                if col >= 2:
                    transp_nome = titulos[col - 1]

                    if transp_nome == "TRANSBUIATTE":
                        num_entregas = len(grupo[['CIDADE_calc', 'CLIENTE_calc']].drop_duplicates())
                        acrescimo = max(0, num_entregas - 1) * 1100

                        cidades_bloco = grupo['CIDADE_calc'].astype(str).str.upper().str.strip().tolist()
                        is_manaus = any('MANAUS' in c for c in cidades_bloco)

                        if is_manaus:
                            if acrescimo > 0: formula_transb = f"=4500+{acrescimo}"
                            else: formula_transb = "=4500"
                        else:
                            if acrescimo > 0: formula_transb = f"=${letra_valor}${linha_total}*0.04+{acrescimo}"
                            else: formula_transb = f"=${letra_valor}${linha_total}*0.04"

                        cel_mesclada.value = formula_transb

                    cel_mesclada.number_format = '"R$" #,##0.00'

            for col in range(1, num_transp + 2):
                cel_tot_esq = ws.cell(row=linha_total, column=col)
                cel_tot_esq.fill = fill_cinza
                cel_tot_esq.font = fonte_preta_normal
                cel_tot_esq.border = borda_fina
                cel_tot_esq.alignment = alinhamento_centro

                if col == 1:
                    cel_tot_esq.value = "% Frete"
                else:
                    letra_col = get_column_letter(col)
                    formula_frete = f"={letra_col}{linha_inicial}/${letra_valor}${linha_total}"
                    cel_tot_esq.value = formula_frete
                    cel_tot_esq.number_format = '0.00%'

            letra_m3 = get_column_letter(col_m3)
            formula_soma_m3 = f"=SUM({letra_m3}{linha_inicial}:{letra_m3}{linha_final})"
            cel_soma_m3 = ws.cell(row=linha_total, column=col_m3, value=formula_soma_m3)
            cel_soma_m3.fill = fill_amarelo
            cel_soma_m3.font = fonte_preta_negrito
            cel_soma_m3.border = borda_fina
            cel_soma_m3.number_format = '0.00'
            cel_soma_m3.alignment = alinhamento_centro

            formula_soma_valor = f"=SUM({letra_valor}{linha_inicial}:{letra_valor}{linha_final})"
            cel_soma_valor = ws.cell(row=linha_total, column=col_valor, value=formula_soma_valor)
            cel_soma_valor.fill = fill_amarelo
            cel_soma_valor.font = fonte_preta_negrito
            cel_soma_valor.border = borda_fina
            cel_soma_valor.number_format = '"R$" #,##0.00'
            cel_soma_valor.alignment = alinhamento_centro

            linha_atual = linha_total + 3

        larguras_padrao = {
            "ITEM": 15, "TRANSBUIATTE": 15, "SPEED": 12, "SR LOG": 12, "BUSSOLA": 12,
            "TODOBRASIL": 15, "AMAZON": 12, "GAB": 15, "BERTOLINI": 18,
            "COSTA BRASIL": 18, "CENTER CARGO": 18, "TECMAR": 18,
            "CIDADE": 20, "UF": 5, "CLIENTE": 45, "CODIGO": 15,
            "QTD": 8, "M³": 10, "Valor": 18, "DESCRIÇÃO": 35
        }
        for col_idx, col_nome in titulos_colunas.items():
            letra = get_column_letter(col_idx)
            ws.column_dimensions[letra].width = larguras_padrao.get(col_nome, 15)

    def gerar_planilha(self):
        arquivo_saida = filedialog.asksaveasfilename(
            title="Onde deseja salvar a Cotação Formatada?",
            defaultextension=".xlsx",
            filetypes=[("Arquivos Excel", "*.xlsx")],
            initialfile="COTACAO_FORMATADA.xlsx"
        )

        if not arquivo_saida: return

        try:
            wb = Workbook()
            wb.remove(wb.active)

            mascara_container = self.df_dados['TIPO_VEICULO_calc'].astype(str).str.contains('Container', case=False, na=False)
            df_cabotagem = self.df_dados[mascara_container]
            df_rodoviario = self.df_dados[~mascara_container]

            if not df_rodoviario.empty:
                ws_rodo = wb.create_sheet(title="Cotação Rodoviário")
                self.gerar_aba_excel(ws_rodo, df_rodoviario, "RODOVIARIO")

            if not df_cabotagem.empty:
                ws_cabo = wb.create_sheet(title="Cotação Cabotagem")
                self.gerar_aba_excel(ws_cabo, df_cabotagem, "CABOTAGEM")

            ws_cif_copy = wb.create_sheet(title="Programação Filtrada (CIF)")
            df_export = self.df_original_cif.copy()
            for col in df_export.columns:
                if str(df_export[col].dtype).startswith('datetime'):
                    df_export[col] = df_export[col].dt.strftime('%d/%m/%Y').fillna('')
                else:
                    df_export[col] = df_export[col].fillna('')

            for r in dataframe_to_rows(df_export, index=False, header=True):
                ws_cif_copy.append(r)

            wb.save(arquivo_saida)
            messagebox.showinfo("Sucesso!", "Planilha Gerada com Sucesso!\n\nAs abas Rodoviário, Cabotagem e Cargas CIF foram salvas perfeitamente.")

        except Exception as e:
            messagebox.showerror("Erro Crítico", f"Ocorreu um erro ao salvar o arquivo Excel:\n\n{str(e)}")

# ==========================================
# CÓDIGO 3: CONSOLIDADOR DE ARQUIVOS (Frame Interno)
# ==========================================
class ConsolidadorApp(ctk.CTkFrame):
    def __init__(self, master=None):
        super().__init__(master, fg_color="transparent")
        self.queue = queue.Queue()
        self.definir_estilos()
        self.criar_widgets()
        self.verificar_queue()

    def definir_estilos(self):
        style = ttk.Style()
        style.theme_use("clam")

        fundo_cinza = "#F4F6F9"
        azul_gree = "#0057B8"

        style.configure(".", background="white", font=("Segoe UI", 10))

        style.configure("TNotebook", background="white", borderwidth=0)
        style.configure("TNotebook.Tab",
                        background="#EAF2FC",
                        foreground="#2C3E50",
                        padding=[2, 3],
                        font=("Segoe UI", 10, "bold"))
        style.map("TNotebook.Tab",
                  background=[("selected", azul_gree)],
                  foreground=[("selected", "#FFFFFF")])

        style.configure("TLabelframe", background="#FFFFFF", bordercolor="#E0E6ED", borderwidth=1)
        style.configure("TLabelframe.Label", font=("Segoe UI", 11, "bold"), foreground=azul_gree, background="#FFFFFF")
        style.configure("TLabel", background="#FFFFFF", foreground="#2C3E50")

        style.configure("TProgressbar", thickness=8, troughcolor="#E0E6ED", background=azul_gree)

    def criar_widgets(self):
        header_frame = tk.Frame(self, bg="#0057B8", height=70)
        header_frame.pack(fill="x", side="top")
        header_frame.pack_propagate(False)

        try:
            logo = ctk.CTkImage(light_image=Image.open(CAMINHO_LOGO), size=(110, 30))
            lblLogo = ctk.CTkLabel(header_frame, image=logo, text="")
            lblLogo.pack(side="left", padx=25, pady=17)
        except Exception:
            pass

        title_text_frame = tk.Frame(header_frame, bg="#0057B8")
        title_text_frame.pack(side="left", fill="both", expand=True, pady=10)

        lbl_brand = tk.Label(
            title_text_frame,
            text="Gree Electric Appliances",
            font=("Segoe UI", 12, "bold"),
            fg="white",
            bg="#0057B8"
        )
        lbl_brand.pack(anchor="w", padx=(5, 15))

        lbl_subbrand = tk.Label(
            title_text_frame,
            text="Painel Integrado de Automação e Gerenciamento de Documentos",
            font=("Segoe UI", 9, "italic"),
            fg="#D0E1F9",
            bg="#0057B8"
        )
        lbl_subbrand.pack(anchor="w", padx=(5, 15))

        body = ctk.CTkFrame(self, fg_color="white", corner_radius=15)
        body.pack(fill="both", expand=True, padx=25, pady=(20, 10))

        self.notebook = ttk.Notebook(body)
        self.notebook.pack(fill="both", expand=True, padx=15, pady=15)

        self.tab_vendas = ctk.CTkFrame(self.notebook, fg_color="white", corner_radius=0)
        self.tab_mesclar = ctk.CTkFrame(self.notebook, fg_color="white", corner_radius=0)
        self.tab_conversor = ctk.CTkFrame(self.notebook, fg_color="white", corner_radius=0)

        self.notebook.add(self.tab_vendas, text=" 📊  Relatório de Ordens de Venda")
        self.notebook.add(self.tab_mesclar, text=" 🔗  Mesclador de PDFs")
        self.notebook.add(self.tab_conversor, text=" ⚙️  Conversor de Arquivos")

        self.montar_aba_vendas()
        self.montar_aba_mesclar()
        self.montar_aba_conversor()

        status_frame = ctk.CTkFrame(self, fg_color="white", corner_radius=12, border_color="#E0E6ED", border_width=1)
        status_frame.pack(fill="x", side="bottom", padx=25, pady=(0, 20), before=body)

        # --- BARRA DE PROGRESSO CORRIGIDA E MODERNA ---
        progress_container = ctk.CTkFrame(status_frame, fg_color="transparent")
        progress_container.pack(fill="x", padx=15, pady=(15, 8))

        self.progress = ctk.CTkProgressBar(progress_container, height=14, progress_color="#2EAF4A", fg_color="#E0E6ED")
        self.progress.pack(side="left", fill="x", expand=True, padx=(0, 15))
        self.progress.set(0)

        self.lbl_porcentagem = ctk.CTkLabel(progress_container, text="0%", font=("Segoe UI", 14, "bold"), text_color="#0057B8")
        self.lbl_porcentagem.pack(side="right")
        # ----------------------------------------------

        self.log_text = tk.Text(
            status_frame,
            height=5,
            font=("Consolas", 9),
            bg="#FFFFFF",
            fg="#2C3E50",
            insertbackground="#0057B8",
            relief="flat",
            borderwidth=0,
            highlightthickness=0
        )
        self.log_text.pack(fill="x", padx=15, pady=(0, 15))
        self.log_text.configure(state="disabled")

        self.escrever_log("Sistema inicializado com sucesso. Aguardando comandos...\n"
                           "Dica: Para evitar travamentos, as operações rodam em segundo plano.")

    def escrever_log(self, mensagem):
        timestamp = dt.now().strftime("[%H:%M:%S]")
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"{timestamp} {mensagem}\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    # --- MATEMÁTICA DA BARRA DE PROGRESSO CORRIGIDA ---
    def verificar_queue(self):
        try:
            while True:
                msg_type, payload = self.queue.get_nowait()
                if msg_type == "log":
                    self.escrever_log(payload)
                elif msg_type == "progress":
                    self.progress.set(payload / 100.0)
                    self.lbl_porcentagem.configure(text=f"{int(payload)}%")
                elif msg_type == "success":
                    self.escrever_log(payload)
                    messagebox.showinfo("Sucesso", payload)
                    self.progress.set(1.0)
                    self.lbl_porcentagem.configure(text="100%")
                elif msg_type == "error":
                    self.escrever_log(f"❌ {payload}")
                    messagebox.showerror("Erro", payload)
                    self.progress.set(0)
                    self.lbl_porcentagem.configure(text="0%")
                self.queue.task_done()
        except queue.Empty:
            pass
        finally:
            self.after(100, self.verificar_queue)
    # ---------------------------------------------------

    def montar_aba_vendas(self):
        container = ctk.CTkFrame(self.tab_vendas, fg_color="white", corner_radius=0)
        container.pack(fill="both", expand=True, padx=15, pady=15)

        mode_group = ctk.CTkFrame(container, fg_color="white", border_color="#E0E6ED", border_width=1, corner_radius=8)
        mode_group.pack(fill="x", pady=(0, 10))

        lbl_mode_title = ctk.CTkLabel(mode_group, text="Modo de Leitura", font=("Segoe UI", 11, "bold"), text_color="#0057B8")
        lbl_mode_title.pack(anchor="w", padx=10, pady=(8, 2))

        inner_mode = ctk.CTkFrame(mode_group, fg_color="transparent")
        inner_mode.pack(fill="x", padx=10, pady=(0, 10))

        self.modo_leitura_var = tk.StringVar(value="pasta")

        r1 = ctk.CTkRadioButton(
            inner_mode,
            text="Pasta de PDFs (Leitura de múltiplos arquivos)",
            variable=self.modo_leitura_var,
            value="pasta",
            fg_color="#0057B8",
            font=("Segoe UI", 11),
            text_color="#2C3E50",
            command=self.atualizar_label_selecao
        )
        r1.pack(side="left", padx=(0, 20), pady=8)

        r2 = ctk.CTkRadioButton(
            inner_mode,
            text="Arquivo PDF Único",
            variable=self.modo_leitura_var,
            value="arquivo",
            fg_color="#0057B8",
            font=("Segoe UI", 11),
            text_color="#2C3E50",
            command=self.atualizar_label_selecao
        )
        r2.pack(side="left", pady=8)

        selection_group = ctk.CTkFrame(container, fg_color="white", border_color="#E0E6ED", border_width=1, corner_radius=8)
        selection_group.pack(fill="x", pady=(0, 10))

        lbl_selection_title = ctk.CTkLabel(selection_group, text="Seleção de Entrada", font=("Segoe UI", 11, "bold"), text_color="#0057B8")
        lbl_selection_title.pack(anchor="w", padx=10, pady=(8, 2))

        inner_selection = ctk.CTkFrame(selection_group, fg_color="transparent")
        inner_selection.pack(fill="x", padx=10, pady=(0, 10))
        inner_selection.columnconfigure(1, weight=1)

        self.lbl_entrada_tipo = ttk.Label(inner_selection, text="Pasta de PDFs:")
        self.lbl_entrada_tipo.grid(row=0, column=0, sticky="w", padx=(0, 5))

        self.ent_vendas_path = ctk.CTkEntry(inner_selection, height=32, corner_radius=6)
        self.ent_vendas_path.grid(row=0, column=1, sticky="ew", padx=(0, 5))

        btn_vendas_proc = ctk.CTkButton(
            inner_selection,
            text="Procurar...",
            command=self.procurar_vendas_path,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#F26522",
            hover_color="#D35400",
            text_color="white",
            font=("Segoe UI", 11, "bold")
        )
        btn_vendas_proc.grid(row=0, column=2, sticky="e")

        config_group = ctk.CTkFrame(container, fg_color="white", border_color="#E0E6ED", border_width=1, corner_radius=8)
        config_group.pack(fill="x", pady=(0, 15))

        lbl_config_title = ctk.CTkLabel(config_group, text="Configurações do Relatório", font=("Segoe UI", 11, "bold"), text_color="#0057B8")
        lbl_config_title.pack(anchor="w", padx=10, pady=(8, 2))

        inner_config = ctk.CTkFrame(config_group, fg_color="transparent")
        inner_config.pack(fill="x", padx=10, pady=(0, 10))

        tk.Label(inner_config, text="Data Doc:", bg="#FFFFFF", font=("Segoe UI", 10)).grid(row=0, column=0, sticky="w", padx=(0, 5), pady=5)
        self.ent_data_doc = ctk.CTkEntry(inner_config, width=120, height=32, corner_radius=6)
        self.ent_data_doc.insert(0, dt.now().strftime("%d/%m/%Y"))
        self.ent_data_doc.grid(row=0, column=1, sticky="w", padx=(0, 20), pady=5)

        tk.Label(inner_config, text="Solicitação:", bg="#FFFFFF", font=("Segoe UI", 10)).grid(row=0, column=2, sticky="w", padx=(0, 5), pady=5)
        self.ent_solicitacao = ctk.CTkEntry(inner_config, width=140, height=32, corner_radius=6)
        self.ent_solicitacao.insert(0, "2026080168")
        self.ent_solicitacao.grid(row=0, column=3, sticky="w", padx=(0, 20), pady=5)





        btn_vendas_gerar = ctk.CTkButton(
            container,
            text="AUTOMATIZAR E GERAR EXCEL",
            command=self.iniciar_vendas,
            height=45,
            corner_radius=8,
            fg_color="#0057B8",
            hover_color="#003A6F",
            font=("Segoe UI", 13, "bold")
        )
        btn_vendas_gerar.pack(fill="x", pady=(5, 0))

    def atualizar_label_selecao(self):
        modo = self.modo_leitura_var.get()
        if modo == "pasta":
            self.lbl_entrada_tipo.configure(text="Pasta de PDFs:")
        else:
            self.lbl_entrada_tipo.configure(text="Arquivo PDF:")

    def procurar_vendas_path(self):
        modo = self.modo_leitura_var.get()
        if modo == "pasta":
            path = filedialog.askdirectory(title="Selecione a pasta de PDFs")
        else:
            path = filedialog.askopenfilename(title="Selecione o arquivo PDF", filetypes=[("Arquivos PDF", "*.pdf")])
        if path:
            self.ent_vendas_path.delete(0, "end")
            self.ent_vendas_path.insert(0, os.path.normpath(path))

    def iniciar_vendas(self):
        entrada = self.ent_vendas_path.get().strip()
        if not entrada:
            messagebox.showerror("Erro", "Selecione uma pasta ou um arquivo PDF de entrada.")
            return
        data_doc = self.ent_data_doc.get().strip()
        solic = self.ent_solicitacao.get().strip()

        self.progress.set(0)
        self.lbl_porcentagem.configure(text="0%")
        self.escrever_log("Iniciando processamento das ordens de venda...")

        # Removemos o itens_col do args.
        thread = threading.Thread(target=self.worker_processar_vendas, args=(entrada, data_doc, solic))
        thread.daemon = True
        thread.start()

    def worker_processar_vendas(self, entrada, data_doc, solic):
        try:
            import pypdf
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.drawing.image import Image as ExcelImage
            import math

            pdf_files = []
            if os.path.isdir(entrada):
                for f in os.listdir(entrada):
                    if f.lower().endswith(".pdf") and not f.startswith("~$") and not f.startswith("Mesclado_"):
                        pdf_files.append(os.path.join(entrada, f))
            elif os.path.isfile(entrada) and entrada.lower().endswith(".pdf"):
                pdf_files.append(entrada)

            if not pdf_files:
                self.queue.put(("error", "Nenhum arquivo PDF válido localizado!"))
                return

            pdf_files.sort()
            self.queue.put(("log", f"Localizado(s) {len(pdf_files)} arquivo(s) PDF para leitura."))
            total_arquivos = len(pdf_files)
            regex_unificado = re.compile(r'\b([S5]\d{8}|100\d{6}|TPZ\d{6}|TPR\d{6})\b', re.IGNORECASE)
            orders_list = []

            for idx, filepath in enumerate(pdf_files):
                filename = os.path.basename(filepath)
                self.queue.put(("log", f"Lendo ({idx+1}/{total_arquivos}): {filename}"))
                self.queue.put(("progress", int(((idx + 1) / total_arquivos) * 90)))
                texto_extraido = ""
                try:
                    with open(filepath, "rb") as f:
                        reader = pypdf.PdfReader(f, strict=False)
                        for page in reader.pages:
                            text = page.extract_text()
                            if text:
                                texto_extraido += text + "\n"
                except Exception:
                    self.queue.put(("log", f"⚠️ Erro ao ler conteúdo interno de {filename}. Usando fallback..."))
                found_any = False
                if texto_extraido:
                    matches = regex_unificado.findall(texto_extraido)
                    if matches:
                        for match in matches:
                            codigo = match.upper()
                            if codigo not in orders_list:
                                orders_list.append(codigo)
                        found_any = True
                if not found_any:
                    nome_sem_ext = os.path.splitext(filename)[0].upper()
                    matches = regex_unificado.findall(nome_sem_ext)
                    if matches:
                        self.queue.put(("log", f"💡 {filename} lido via nome do arquivo (fallback)."))
                        for match in matches:
                            codigo = match.upper()
                            # Manter regra caso exista a variavel padronizar no topo do seu script
                            try:
                                if padronizar and codigo.startswith("5"):
                                    codigo = "S" + codigo[1:]
                            except:
                                pass
                            if codigo not in orders_list:
                                orders_list.append(codigo)
                    else:
                        self.queue.put(("log", f"⚠️ Nenhum padrão encontrado em {filename}."))
                gc.collect()

            if not orders_list:
                self.queue.put(("error", "Nenhum código encontrado nos arquivos!"))
                return

            self.queue.put(("log", f"Total de {len(orders_list)} ordens capturadas! Organizando planilha..."))
            orders_ordenados = orders_list

            # === NOVO CÁLCULO AUTOMÁTICO DE LINHAS ===
            total_pedidos = len(orders_ordenados)
            itens_col = math.ceil(total_pedidos / 6)
            if itens_col == 0:
                itens_col = 1

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Controle de Entrega"
            ws.views.sheetView[0].showGridLines = True

            font_title = Font(name="Segoe UI", size=16, bold=True, color="1F4E78")
            font_subtitle = Font(name="Segoe UI", size=12, bold=True, color="595959")
            font_meta = Font(name="Segoe UI", size=10, bold=True, color="000000")
            font_meta_center = Font(name="Segoe UI", size=10, bold=True, color="595959")
            font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")
            font_data_reg = Font(name="Segoe UI", size=10, bold=False, color="000000")
            font_footer_label = Font(name="Segoe UI", size=10, bold=True, color="2F5496")
            font_footer_line = Font(name="Segoe UI", size=11, bold=False, color="595959")
            fill_header = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
            fill_zebra_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            fill_zebra_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            fill_special = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            thin_side = Side(border_style="thin", color="D9D9D9")
            border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

            ws.column_dimensions['A'].width = 3
            for col_idx in range(6):
                ped_col = chr(ord('B') + col_idx * 2)
                ok_col = chr(ord('C') + col_idx * 2)
                ws.column_dimensions[ped_col].width = 14
                ws.column_dimensions[ok_col].width = 5

            # === NOVA LÓGICA DE INSERÇÃO DA LOGO ===
            ws.merge_cells("B2:M2")
            ws.row_dimensions[2].height = 45

            try:
                img_logo = ExcelImage(CAMINHO_LOGO_COR)
                ws.add_image(img_logo, "B2")
            except Exception as e:
                self.queue.put(("log", f"Aviso sobre a logo: {str(e)}"))
                ws["B2"] = "GREE ELECTRIC APPLIANCES DO BRASIL LTDA."
                ws["B2"].font = font_title
                ws["B2"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("B4:M4")
            ws["B4"] = "Controle de Entrega - Pedidos de Vendas"
            ws["B4"].font = font_subtitle
            ws["B4"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("B7:D7")
            ws["B7"] = f"Data do Documento: {data_doc}"
            ws["B7"].font = font_meta
            ws["B7"].alignment = Alignment(horizontal="left", vertical="center")

            ws.merge_cells("G7:I7")
            ws["G7"] = "SISTEMA GREE APP"
            ws["G7"].font = font_meta_center
            ws["G7"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("K7:M7")
            ws["K7"] = f"Solicitação: {solic}"
            ws["K7"].font = font_meta
            ws["K7"].alignment = Alignment(horizontal="right", vertical="center")

            for col_idx in range(6):
                ped_col = chr(ord('B') + col_idx * 2)
                ok_col = chr(ord('C') + col_idx * 2)
                ws[f"{ped_col}9"] = "Pedido"
                ws[f"{ped_col}9"].font = font_header
                ws[f"{ped_col}9"].fill = fill_header
                ws[f"{ped_col}9"].alignment = Alignment(horizontal="center", vertical="center")
                ws[f"{ped_col}9"].border = border_cell
                ws[f"{ok_col}9"] = "OK"
                ws[f"{ok_col}9"].font = font_header
                ws[f"{ok_col}9"].fill = fill_header
                ws[f"{ok_col}9"].alignment = Alignment(horizontal="center", vertical="center")
                ws[f"{ok_col}9"].border = border_cell

            total_slots = itens_col * 6
            for i in range(total_slots):
                col_num = i // itens_col
                row_num = 10 + (i % itens_col)
                ped_col = chr(ord('B') + col_num * 2)
                ok_col = chr(ord('C') + col_num * 2)

                order_val = orders_ordenados[i] if i < len(orders_ordenados) else ""
                is_special = order_val.startswith("100") or order_val.startswith("TPZ") or order_val.startswith("TPR")
                is_even_row = (row_num % 2 == 0)
                row_fill = fill_zebra_white if is_even_row else fill_zebra_gray

                if order_val:
                    if is_special:
                        ws[f"{ped_col}{row_num}"] = order_val
                        ws[f"{ped_col}{row_num}"].font = font_data_bold
                        ws[f"{ped_col}{row_num}"].fill = fill_special
                        ws[f"{ok_col}{row_num}"] = "☐"
                        ws[f"{ok_col}{row_num}"].font = font_data_bold
                        ws[f"{ok_col}{row_num}"].fill = fill_special
                    else:
                        ws[f"{ped_col}{row_num}"] = order_val
                        ws[f"{ped_col}{row_num}"].font = font_data_reg
                        ws[f"{ped_col}{row_num}"].fill = row_fill
                        ws[f"{ok_col}{row_num}"] = "☐"
                        ws[f"{ok_col}{row_num}"].font = font_data_reg
                        ws[f"{ok_col}{row_num}"].fill = row_fill
                else:
                    ws[f"{ped_col}{row_num}"] = ""
                    ws[f"{ped_col}{row_num}"].font = font_data_reg
                    ws[f"{ped_col}{row_num}"].fill = row_fill
                    ws[f"{ok_col}{row_num}"] = "☐"
                    ws[f"{ok_col}{row_num}"].font = font_data_reg
                    ws[f"{ok_col}{row_num}"].fill = row_fill

                ws[f"{ped_col}{row_num}"].alignment = Alignment(horizontal="left", vertical="center")
                ws[f"{ped_col}{row_num}"].border = border_cell
                ws[f"{ok_col}{row_num}"].alignment = Alignment(horizontal="center", vertical="center")
                ws[f"{ok_col}{row_num}"].border = border_cell

            footer_row = 10 + itens_col + 2
            ws.merge_cells(f"B{footer_row}:E{footer_row}")
            ws[f"B{footer_row}"] = "________________________________________________"
            ws[f"B{footer_row}"].font = font_footer_line
            ws[f"B{footer_row}"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"I{footer_row}:L{footer_row}")
            ws[f"I{footer_row}"] = "______ / ______ / 2026"
            ws[f"I{footer_row}"].font = font_footer_line
            ws[f"I{footer_row}"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"B{footer_row+1}:E{footer_row+1}")
            ws[f"B{footer_row+1}"] = "Responsável"
            ws[f"B{footer_row+1}"].font = font_footer_label
            ws[f"B{footer_row+1}"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"I{footer_row+1}:L{footer_row+1}")
            ws[f"I{footer_row+1}"] = "Data de Recebimento"
            ws[f"I{footer_row+1}"].font = font_footer_label
            ws[f"I{footer_row+1}"].alignment = Alignment(horizontal="center", vertical="center")

            diretorio_destino = entrada if os.path.isdir(entrada) else os.path.dirname(entrada)
            filename_output = f"controle_de_entrega_venda_Gree_{dt.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            out_xlsx = os.path.join(diretorio_destino, filename_output)

            wb.save(out_xlsx)
            wb.close()
            gc.collect()

            self.queue.put(("progress", 100))
            self.queue.put(("success", f"Relatório gerado com sucesso!\nSalvo em: {out_xlsx}"))

        except Exception as e:
            self.queue.put(("error", f"Ocorreu um erro no processamento das ordens: {str(e)}"))

    def montar_aba_mesclar(self):
        container = ctk.CTkFrame(self.tab_mesclar, fg_color="white", corner_radius=0)
        container.pack(fill="both", expand=True, padx=15, pady=15)

        lbl_info = ctk.CTkLabel(container, text="Junte múltiplos PDFs de faturamento, Notas Fiscais ou laudos em um único arquivo consolidado.", font=("Segoe UI", 11, "italic"), text_color="gray40")
        lbl_info.pack(anchor="w", pady=(0, 10))

        list_frame = ctk.CTkFrame(container, fg_color="transparent")
        list_frame.pack(fill="both", expand=True, pady=(0, 10))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)

        frame_lista = ctk.CTkFrame(list_frame, fg_color="white", border_color="#BDC3C7", border_width=1, corner_radius=8)
        frame_lista.grid(row=0, column=0, sticky="nsew")

        self.lst_pdfs = tk.Listbox(
            frame_lista,
            font=("Segoe UI", 11),
            selectbackground="#0057B8",
            selectforeground="white",
            relief="flat",
            borderwidth=0,
            highlightthickness=0,
            bg="white",
            fg="#2C3E50"
        )
        self.lst_pdfs.pack(side="left", fill="both", expand=True, padx=8, pady=8)

        scrollbar = ttk.Scrollbar(frame_lista, orient="vertical", command=self.lst_pdfs.yview)
        scrollbar.pack(side="right", fill="y", padx=(0, 4), pady=4)
        self.lst_pdfs.configure(yscrollcommand=scrollbar.set)

        btn_control_frame = ctk.CTkFrame(list_frame, fg_color="transparent")
        btn_control_frame.grid(row=0, column=2, sticky="ns", padx=(10, 0))

        btn_up = ctk.CTkButton(
            btn_control_frame,
            text="Mover ↑",
            command=self.mover_item_cima,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#0057B8",
            text_color="white",
            hover_color="#003A6F",
            font=("Segoe UI", 10, "bold")
        )
        btn_up.pack(fill="x", pady=2)

        btn_down = ctk.CTkButton(
            btn_control_frame,
            text="Mover ↓",
            command=self.mover_item_baixo,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#0057B8",
            text_color="white",
            hover_color="#003A6F",
            font=("Segoe UI", 10, "bold")
        )
        btn_down.pack(fill="x", pady=2)

        btn_rem = ctk.CTkButton(
            btn_control_frame,
            text="Remover",
            command=self.remover_item,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#E74C3C",
            text_color="white",
            hover_color="#C0392B",
            font=("Segoe UI", 10, "bold")
        )
        btn_rem.pack(fill="x", pady=15)

        btn_clean = ctk.CTkButton(
            btn_control_frame,
            text="Limpar Lista",
            command=self.limpar_lista,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#7F8C8D",
            text_color="white",
            hover_color="#626567",
            font=("Segoe UI", 10, "bold")
        )
        btn_clean.pack(fill="x", pady=2)

        # ==========================================
        # BARRA DE AÇÕES (Importação e Mesclagem)
        # ==========================================
        action_bar = ctk.CTkFrame(container, fg_color="transparent")
        action_bar.pack(fill="x", side="bottom", pady=(10, 0))

        # O segredo: a coluna 2 será uma "mola invisível" que empurra os lados
        action_bar.columnconfigure(2, weight=1)

        btn_add_files = ctk.CTkButton(
            action_bar,
            text="+ Adicionar PDFs individuais...",
            command=self.adicionar_pdfs_individuais,
            height=36,
            corner_radius=8,
            fg_color="#F26522",
            text_color="white",
            hover_color="#D35400",
            font=("Segoe UI", 11, "bold")
        )
        # Travado na coluna 0, alinhado à esquerda (sticky="w")
        btn_add_files.grid(row=0, column=0, sticky="w", padx=(0, 10))

        btn_add_dir = ctk.CTkButton(
            action_bar,
            text="+ Importar de uma pasta...",
            command=self.adicionar_pdfs_pasta,
            height=36,
            corner_radius=8,
            fg_color="#F26522",
            text_color="white",
            hover_color="#D35400",
            font=("Segoe UI", 11, "bold")
        )
        # Travado na coluna 1, alinhado à esquerda (sticky="w")
        btn_add_dir.grid(row=0, column=1, sticky="w")

        self.btn_consolidar = ctk.CTkButton(
            action_bar,
            text="MESCLAR E SALVAR PDF ÚNICO",
            command=self.iniciar_mesclagem,
            height=36,
            width=450,  # <-- PODE MUDAR A LARGURA À VONTADE AQUI
            corner_radius=8,
            fg_color="#2EAF4A",
            hover_color="#22953D",
            font=("Segoe UI", 12, "bold")
        )
        # Travado na coluna 3, alinhado totalmente à direita (sticky="e")
        self.btn_consolidar.grid(row=0, column=3, sticky="e")

        self.arquivos_mescla_lista = []

    def adicionar_pdfs_individuais(self):
        files = filedialog.askopenfilenames(title="Selecione arquivos PDF", filetypes=[("Arquivos PDF", "*.pdf")])
        if files:
            for f in files:
                f_normalized = os.path.normpath(f)
                if f_normalized not in self.arquivos_mescla_lista:
                    self.arquivos_mescla_lista.append(f_normalized)
                    self.lst_pdfs.insert("end", os.path.basename(f_normalized))
            self.escrever_log(f"Adicionados {len(files)} arquivo(s) à lista.")

    # --- FUNÇÃO DE VARREDURA RECURSIVA ---
    def adicionar_pdfs_pasta(self):
        folder = filedialog.askdirectory(title="Selecione a pasta principal (Lê subpastas automaticamente)")
        if folder:
            encontrados = 0
            # os.walk vasculha a pasta selecionada e todas as subpastas dentro dela
            for root, dirs, files in os.walk(folder):
                for f in sorted(files):
                    if f.lower().endswith(".pdf") and not f.startswith("~$") and not f.startswith("Mesclado_"):
                        full_path = os.path.normpath(os.path.join(root, f))
                        if full_path not in self.arquivos_mescla_lista:
                            self.arquivos_mescla_lista.append(full_path)
                            self.lst_pdfs.insert("end", f)
                            encontrados += 1
            if encontrados > 0:
                self.escrever_log(f"Adicionados {encontrados} arquivo(s) PDF importados da pasta e subpastas.")
            else:
                messagebox.showinfo("Informação", "Nenhum arquivo PDF válido localizado nessa pasta ou subpastas!")
    # -------------------------------------

    def mover_item_cima(self):
        sel = self.lst_pdfs.curselection()
        if not sel: return
        idx = sel[0]
        if idx == 0: return
        self.arquivos_mescla_lista[idx], self.arquivos_mescla_lista[idx - 1] = self.arquivos_mescla_lista[idx - 1], self.arquivos_mescla_lista[idx]
        text = self.lst_pdfs.get(idx)
        self.lst_pdfs.delete(idx)
        self.lst_pdfs.insert(idx - 1, text)
        self.lst_pdfs.selection_set(idx - 1)

    def mover_item_baixo(self):
        sel = self.lst_pdfs.curselection()
        if not sel: return
        idx = sel[0]
        if idx == self.lst_pdfs.size() - 1: return
        self.arquivos_mescla_lista[idx], self.arquivos_mescla_lista[idx + 1] = self.arquivos_mescla_lista[idx + 1], self.arquivos_mescla_lista[idx]
        text = self.lst_pdfs.get(idx)
        self.lst_pdfs.delete(idx)
        self.lst_pdfs.insert(idx + 1, text)
        self.lst_pdfs.selection_set(idx + 1)

    def remover_item(self):
        sel = self.lst_pdfs.curselection()
        if not sel: return
        idx = sel[0]
        self.lst_pdfs.delete(idx)
        self.arquivos_mescla_lista.pop(idx)

    def limpar_lista(self):
        self.lst_pdfs.delete(0, "end")
        self.arquivos_mescla_lista.clear()
        self.escrever_log("Lista limpa pelo usuário.")

    def iniciar_mesclagem(self):
        if not self.arquivos_mescla_lista:
            messagebox.showerror("Erro", "Sua lista de arquivos está vazia!")
            return
        save_path = filedialog.asksaveasfilename(
            title="Salvar PDF Mesclado Como",
            defaultextension=".pdf", filetypes=[("Documento PDF", "*.pdf")],
            initialfile=f"Mesclado_Ordens_Gree_{dt.now().strftime('%Y%m%d')}.pdf"
        )
        if save_path:
            save_path = os.path.normpath(save_path)
            self.progress.set(0)
            self.lbl_porcentagem.configure(text="0%")
            self.escrever_log("Iniciando processo de mesclagem...")
            thread = threading.Thread(target=self.worker_mesclar_pdfs, args=(save_path,))
            thread.daemon = True
            thread.start()

    def worker_mesclar_pdfs(self, dest_path):
        try:
            import pypdf
            writer = pypdf.PdfWriter()
            total_arquivos = len(self.arquivos_mescla_lista)
            arquivos_anexados = 0
            for idx, file_path in enumerate(self.arquivos_mescla_lista):
                filename = os.path.basename(file_path)
                if os.path.normpath(file_path) == os.path.normpath(dest_path):
                    continue
                self.queue.put(("log", f"Anexando PDF ({idx+1}/{total_arquivos}): {filename}"))
                self.queue.put(("progress", int(((idx + 1) / total_arquivos) * 90)))
                try:
                    with open(file_path, "rb") as f_in:
                        reader = pypdf.PdfReader(f_in, strict=False)
                        for page in reader.pages:
                            writer.add_page(page)
                        arquivos_anexados += 1
                except Exception as ex:
                    self.queue.put(("log", f"⚠️ Arquivo ignorado: {filename} ({str(ex)})"))
            if arquivos_anexados == 0:
                self.queue.put(("error", "Nenhum arquivo válido pôde ser lido!"))
                return
            self.queue.put(("log", "Gravando arquivo unificado..."))
            with open(dest_path, "wb") as f_out:
                writer.write(f_out)
            writer.close()
            gc.collect()
            self.queue.put(("progress", 100))
            self.queue.put(("success", f"Mesclagem finalizada!\nArquivo gravado: {dest_path}"))
        except Exception as e:
            self.queue.put(("error", f"Erro ao mesclar: {str(e)}"))

    def montar_aba_conversor(self):
        container = ctk.CTkFrame(self.tab_conversor, fg_color="white", corner_radius=0)
        container.pack(fill="both", expand=True, padx=15, pady=15)

        lbl_info = ctk.CTkLabel(container, text="Selecione abaixo a ferramenta de conversão rápida que deseja utilizar.", font=("Segoe UI", 11, "italic"), text_color="gray40")
        lbl_info.pack(anchor="w", pady=(0, 10))

        self.sub_notebook = ttk.Notebook(container)
        self.sub_notebook.pack(fill="both", expand=True)

        self.tab_pdf_word = ctk.CTkFrame(self.sub_notebook, fg_color="white", corner_radius=0)
        self.tab_pdf_jpg = ctk.CTkFrame(self.sub_notebook, fg_color="white", corner_radius=0)
        self.tab_jpg_pdf = ctk.CTkFrame(self.sub_notebook, fg_color="white", corner_radius=0)

        self.sub_notebook.add(self.tab_pdf_word, text="PDF para Word")
        self.sub_notebook.add(self.tab_pdf_jpg, text="PDF para JPG")
        self.sub_notebook.add(self.tab_jpg_pdf, text="JPG para PDF")

        self.montar_sub_pdf_word()
        self.montar_sub_pdf_jpg()
        self.montar_sub_jpg_pdf()

    def montar_sub_pdf_word(self):
        container = ctk.CTkFrame(self.tab_pdf_word, fg_color="white", corner_radius=0)
        container.pack(fill="both", expand=True, padx=15, pady=15)

        lbl_desc = ctk.CTkLabel(container, text="Converta arquivos PDF em documentos do Word (.docx) editáveis.", font=("Segoe UI", 11, "italic"), text_color="gray40")
        lbl_desc.pack(anchor="w", pady=(0, 15))

        selection_group = ctk.CTkFrame(container, fg_color="white", border_color="#E0E6ED", border_width=1, corner_radius=8)
        selection_group.pack(fill="x", pady=(0, 20))

        lbl_sel_title = ctk.CTkLabel(selection_group, text="Seleção de Arquivo", font=("Segoe UI", 11, "bold"), text_color="#0057B8")
        lbl_sel_title.pack(anchor="w", padx=10, pady=(8, 2))

        inner_selection = ctk.CTkFrame(selection_group, fg_color="transparent")
        inner_selection.pack(fill="x", padx=10, pady=(0, 10))
        inner_selection.columnconfigure(1, weight=1)

        ctk.CTkLabel(inner_selection, text="Arquivo PDF:", font=("Segoe UI", 11), text_color="#2C3E50").grid(row=0, column=0, sticky="w", padx=(0, 5))

        self.ent_pdf_word_path = ctk.CTkEntry(inner_selection, height=32, corner_radius=6)
        self.ent_pdf_word_path.grid(row=0, column=1, sticky="ew", padx=(0, 5))

        btn_proc = ctk.CTkButton(
            inner_selection,
            text="Procurar...",
            command=self.procurar_pdf_word,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#F26522",
            hover_color="#D35400",
            text_color="white",
            font=("Segoe UI", 11, "bold")
        )
        btn_proc.grid(row=0, column=2, sticky="e")

        btn_gerar = ctk.CTkButton(
            container,
            text="CONVERTER PARA WORD",
            command=self.iniciar_pdf_word,
            height=45,
            corner_radius=8,
            fg_color="#0057B8",
            hover_color="#003A6F",
            font=("Segoe UI", 13, "bold")
        )
        btn_gerar.pack(fill="x", pady=(5, 0))

    def procurar_pdf_word(self):
        path = filedialog.askopenfilename(title="Selecione o arquivo PDF", filetypes=[("Arquivos PDF", "*.pdf")])
        if path:
            self.ent_pdf_word_path.delete(0, "end")
            self.ent_pdf_word_path.insert(0, os.path.normpath(path))

    def iniciar_pdf_word(self):
        path = self.ent_pdf_word_path.get().strip()
        if not path:
            messagebox.showerror("Erro", "Selecione um arquivo de entrada!")
            return
        self.progress.set(0)
        self.lbl_porcentagem.configure(text="0%")
        self.escrever_log("Iniciando conversão para Word...")
        thread = threading.Thread(target=self.worker_convert_pdf_to_word, args=(path,))
        thread.daemon = True
        thread.start()

    def worker_convert_pdf_to_word(self, pdf_path):
        try:
            import pypdf
            import docx
            from docx import Document
            import tempfile

            pypdfium_available = False
            try:
                import pypdfium2 as pdfium
                pypdfium_available = True
            except ImportError:
                pass

            self.queue.put(("log", f"Iniciando conversão de PDF para Word..."))
            filename = os.path.basename(pdf_path)
            dir_name = os.path.dirname(pdf_path)
            base_name = os.path.splitext(filename)[0]
            out_docx = os.path.join(dir_name, f"{base_name}.docx")

            try:
                if os.path.exists(out_docx):
                    with open(out_docx, "a+b") as f: pass
            except IOError:
                self.queue.put(("error", "O arquivo Word de destino já está aberto!"))
                return

            doc = Document()
            doc.add_heading(base_name, level=0)

            with open(pdf_path, "rb") as f:
                reader = pypdf.PdfReader(f, strict=False)
                total_pages = len(reader.pages)
                has_any_text = False
                for page in reader.pages:
                    t = page.extract_text()
                    if t and t.strip():
                        has_any_text = True
                        break

                if has_any_text:
                    self.queue.put(("log", "Extraindo parágrafos..."))
                    for i, page in enumerate(reader.pages):
                        self.queue.put(("log", f"Lendo página ({i+1}/{total_pages})..."))
                        self.queue.put(("progress", int(((i + 1) / total_pages) * 90)))
                        text = page.extract_text()
                        if text:
                            for paragraph in text.split("\n"):
                                if paragraph.strip():
                                    doc.add_paragraph(paragraph)
                        if i < total_pages - 1:
                            doc.add_page_break()
                else:
                    if pypdfium_available:
                        self.queue.put(("log", "Convertendo páginas escaneadas..."))
                        pdf_doc = pdfium.PdfDocument(pdf_path)
                        for i in range(total_pages):
                            self.queue.put(("log", f"Processando ({i+1}/{total_pages})..."))
                            self.queue.put(("progress", int(((i + 1) / total_pages) * 90)))
                            page = pdf_doc[i]
                            bitmap = page.render(scale=2.0)
                            pil_img = bitmap.to_pil()
                            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp_file:
                                tmp_path = tmp_file.name
                            pil_img.save(tmp_path, "JPEG")
                            doc.add_picture(tmp_path, width=docx.shared.Inches(5.8))
                            try: os.remove(tmp_path)
                            except: pass
                            if i < total_pages - 1:
                                doc.add_page_break()
                        pdf_doc.close()
                    else:
                        doc.add_paragraph("Este é um PDF escaneado (imagem). pypdfium2 necessário para imagens.")

            self.queue.put(("log", "Gravando documento..."))
            doc.save(out_docx)
            self.queue.put(("progress", 100))
            self.queue.put(("success", f"PDF convertido com sucesso!\nSalvo em: {out_docx}"))
        except Exception as e:
            self.queue.put(("error", f"Erro na conversão: {str(e)}"))

    def montar_sub_pdf_jpg(self):
        container = ctk.CTkFrame(self.tab_pdf_jpg, fg_color="white", corner_radius=0)
        container.pack(fill="both", expand=True, padx=15, pady=15)

        lbl_desc = ctk.CTkLabel(container, text="Extraia as páginas do PDF para imagens JPG.", font=("Segoe UI", 11, "italic"), text_color="gray40")
        lbl_desc.pack(anchor="w", pady=(0, 15))

        config_group = ctk.CTkFrame(container, fg_color="white", border_color="#E0E6ED", border_width=1, corner_radius=8)
        config_group.pack(fill="x", pady=(0, 20))

        lbl_conf_title = ctk.CTkLabel(config_group, text="Seleção de Arquivo e Opções", font=("Segoe UI", 11, "bold"), text_color="#0057B8")
        lbl_conf_title.pack(anchor="w", padx=10, pady=(8, 2))

        inner_config = ctk.CTkFrame(config_group, fg_color="transparent")
        inner_config.pack(fill="x", padx=10, pady=(0, 10))
        inner_config.columnconfigure(1, weight=1)

        ctk.CTkLabel(inner_config, text="Arquivo PDF:", font=("Segoe UI", 11), text_color="#2C3E50").grid(row=0, column=0, sticky="w", padx=(0, 5), pady=5)

        self.ent_pdf_jpg_path = ctk.CTkEntry(inner_config, height=32, corner_radius=6)
        self.ent_pdf_jpg_path.grid(row=0, column=1, sticky="ew", padx=(0, 5), pady=5)

        btn_proc = ctk.CTkButton(
            inner_config,
            text="Procurar...",
            command=self.procurar_pdf_jpg,
            width=100,
            height=32,
            corner_radius=6,
            fg_color="#F26522",
            hover_color="#D35400",
            text_color="white",
            font=("Segoe UI", 11, "bold")
        )
        btn_proc.grid(row=0, column=2, sticky="e", pady=5)

        ctk.CTkLabel(inner_config, text="Qualidade (DPI):", font=("Segoe UI", 11), text_color="#2C3E50").grid(row=1, column=0, sticky="w", padx=(0, 5), pady=5)
        self.cmb_dpi = ttk.Combobox(inner_config, values=["150 (Médio)", "200 (Bom)", "300 (Alto)"], state="readonly", width=15)
        self.cmb_dpi.set("150 (Médio)")
        self.cmb_dpi.grid(row=1, column=1, sticky="w", pady=5)

        btn_gerar = ctk.CTkButton(
            container,
            text="CONVERTER PARA JPG",
            command=self.iniciar_pdf_jpg,
            height=45,
            corner_radius=8,
            fg_color="#2EAF4A",
            hover_color="#22953D",
            font=("Segoe UI", 13, "bold")
        )
        btn_gerar.pack(fill="x", pady=(5, 0))

    def procurar_pdf_jpg(self):
        path = filedialog.askopenfilename(title="Selecione o arquivo PDF", filetypes=[("Arquivos PDF", "*.pdf")])
        if path:
            self.ent_pdf_jpg_path.delete(0, "end")
            self.ent_pdf_jpg_path.insert(0, os.path.normpath(path))

    def iniciar_pdf_jpg(self):
        path = self.ent_pdf_jpg_path.get().strip()
        if not path:
            messagebox.showerror("Erro", "Selecione um arquivo de entrada!")
            return
        dpi_str = self.cmb_dpi.get().split()[0]
        try: dpi = int(dpi_str)
        except ValueError: dpi = 150
        self.progress.set(0)
        self.lbl_porcentagem.configure(text="0%")
        self.escrever_log("Iniciando extração para JPG...")
        thread = threading.Thread(target=self.worker_convert_pdf_to_jpg, args=(path, dpi))
        thread.daemon = True
        thread.start()

    def worker_convert_pdf_to_jpg(self, pdf_path, dpi):
        try:
            import pypdfium2 as pdfium

            self.queue.put(("log", f"Extraindo imagens via pypdfium2..."))
            filename = os.path.basename(pdf_path)
            dir_name = os.path.dirname(pdf_path)
            base_name = os.path.splitext(filename)[0]

            out_dir = os.path.join(dir_name, f"Imagens_{base_name}")
            os.makedirs(out_dir, exist_ok=True)

            doc = pdfium.PdfDocument(pdf_path)
            total_pages = len(doc)
            zoom = dpi / 72.0

            for i in range(total_pages):
                self.queue.put(("log", f"Renderizando {i+1}/{total_pages}..."))
                self.queue.put(("progress", int(((i + 1) / total_pages) * 90)))
                page = doc[i]
                bitmap = page.render(scale=zoom)
                pil_img = bitmap.to_pil()
                out_file = os.path.join(out_dir, f"pagina_{i+1:03d}.jpg")
                pil_img.save(out_file, "JPEG")

            doc.close()
            self.queue.put(("progress", 100))
            self.queue.put(("success", f"Convertido com sucesso!\nPasta: {out_dir}"))
        except Exception as e:
            self.queue.put(("error", f"Erro: {str(e)}"))

    def montar_sub_jpg_pdf(self):
        container = ctk.CTkFrame(self.tab_jpg_pdf, fg_color="white", corner_radius=0)
        container.pack(fill="both", expand=True, padx=15, pady=10)

        lbl_desc = ctk.CTkLabel(container, text="Una fotos em um único arquivo PDF.", font=("Segoe UI", 11, "italic"), text_color="gray40")
        lbl_desc.pack(anchor="w", pady=(0, 5))

        # 1. CONFIGURAÇÕES COM O BOTÃO INCLUÍDO
        config_group = ctk.CTkFrame(container, fg_color="white", border_color="#E0E6ED", border_width=1, corner_radius=8)
        config_group.pack(side="bottom", fill="x", pady=(0, 0))

        lbl_layout_title = ctk.CTkLabel(config_group, text="Configurações de Layout", font=("Segoe UI", 11, "bold"), text_color="#0057B8")
        lbl_layout_title.pack(anchor="w", padx=10, pady=(8, 2))

        inner_config = ctk.CTkFrame(config_group, fg_color="transparent")
        inner_config.pack(fill="x", padx=10, pady=(0, 10))

        # A coluna 4 funcionará como uma "mola", empurrando o botão para a direita
        inner_config.columnconfigure(4, weight=1)

        ctk.CTkLabel(inner_config, text="Orientação:", font=("Segoe UI", 11), text_color="#2C3E50").grid(row=0, column=0, sticky="w", padx=(0, 5))
        self.cmb_orientacao = ttk.Combobox(inner_config, values=["Retrato (Vertical)", "Paisagem (Horizontal)"], state="readonly", width=18)
        self.cmb_orientacao.set("Retrato (Vertical)")
        self.cmb_orientacao.grid(row=0, column=1, sticky="w", padx=(0, 15))

        ctk.CTkLabel(inner_config, text="Margem:", font=("Segoe UI", 11), text_color="#2C3E50").grid(row=0, column=2, sticky="w", padx=(0, 5))
        self.cmb_margem = ttk.Combobox(inner_config, values=["Sem margem", "Margem fina", "Margem larga"], state="readonly", width=15)
        self.cmb_margem.set("Sem margem")
        self.cmb_margem.grid(row=0, column=3, sticky="w")

        # Novo botão posicionado na extremidade direita (coluna 5)
        btn_gerar = ctk.CTkButton(
            inner_config,
            text="GERAR PDF",
            command=self.iniciar_jpg_pdf,
            height=36,
            width=200,
            corner_radius=8,
            fg_color="#2EAF4A",
            hover_color="#22953D",
            font=("Segoe UI", 12, "bold")
        )
        btn_gerar.grid(row=0, column=5, sticky="e", padx=(10, 0))

        # 2. LISTA NO CENTRO DA TELA
        list_frame = ctk.CTkFrame(container, fg_color="transparent")
        list_frame.pack(side="top", fill="both", expand=True, pady=(0, 15))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)

        frame_lista = ctk.CTkFrame(list_frame, fg_color="white", border_color="#BDC3C7", border_width=1, corner_radius=8)
        frame_lista.grid(row=0, column=0, sticky="nsew")

        self.lst_jpgs = tk.Listbox(
            frame_lista,
            font=("Segoe UI", 11),
            selectbackground="#0057B8",
            selectforeground="white",
            relief="flat",
            borderwidth=0,
            highlightthickness=0,
            bg="white",
            fg="#2C3E50"
        )
        self.lst_jpgs.pack(side="left", fill="both", expand=True, padx=8, pady=8)

        scrollbar = ttk.Scrollbar(frame_lista, orient="vertical", command=self.lst_jpgs.yview)
        scrollbar.pack(side="right", fill="y", padx=(0, 4), pady=4)
        self.lst_jpgs.configure(yscrollcommand=scrollbar.set)

        # 3. BARRA LATERAL DESCOMPRIMIDA
        btn_control_frame = ctk.CTkFrame(list_frame, fg_color="transparent")
        btn_control_frame.grid(row=0, column=2, sticky="ns", padx=(10, 0))

        # Agora os botões laterais têm espaço de sobra
        btn_add = ctk.CTkButton(btn_control_frame, text="Adicionar...", command=self.adicionar_imagens_individuais, width=100, height=25, corner_radius=6, fg_color="#F26522", text_color="white", hover_color="#D35400", font=("Segoe UI", 10, "bold"))
        btn_add.pack(fill="x", pady=2)

        btn_up = ctk.CTkButton(btn_control_frame, text="Mover ↑", command=self.mover_jpg_cima, width=100, height=20, corner_radius=6, fg_color="#0057B8", text_color="white", hover_color="#003A6F", font=("Segoe UI", 10, "bold"))
        btn_up.pack(fill="x", pady=2)

        btn_down = ctk.CTkButton(btn_control_frame, text="Mover ↓", command=self.mover_jpg_baixo, width=100, height=20, corner_radius=6, fg_color="#0057B8", text_color="white", hover_color="#003A6F", font=("Segoe UI", 10, "bold"))
        btn_down.pack(fill="x", pady=2)

        btn_rem = ctk.CTkButton(btn_control_frame, text="Remover", command=self.remover_jpg_item, width=100, height=20, corner_radius=6, fg_color="#E74C3C", text_color="white", hover_color="#C0392B", font=("Segoe UI", 10, "bold"))
        btn_rem.pack(fill="x", pady=2)

        btn_clean = ctk.CTkButton(btn_control_frame, text="Limpar", command=self.limpar_jpg_lista, width=100, height=20, corner_radius=6, fg_color="#7F8C8D", text_color="white", hover_color="#626567", font=("Segoe UI", 10, "bold"))
        btn_clean.pack(fill="x", pady=2)

        self.arquivos_jpg_lista = []


        # ==========================================
        # BOTÃO FINAL DE AÇÃO
        # ==========================================
        btn_gerar = ctk.CTkButton(
            container,
            text="GERAR PDF",
            command=self.iniciar_jpg_pdf,
            height=45,
            corner_radius=8,
            fg_color="#2EAF4A",
            hover_color="#22953D",
            font=("Segoe UI", 13, "bold")
        )
        btn_gerar.pack(fill="x", pady=(15, 0))

    def adicionar_imagens_individuais(self):
        files = filedialog.askopenfilenames(filetypes=[("Imagens", "*.jpg;*.jpeg;*.png")])
        if files:
            for f in files:
                f_normalized = os.path.normpath(f)
                if f_normalized not in self.arquivos_jpg_lista:
                    self.arquivos_jpg_lista.append(f_normalized)
                    self.lst_jpgs.insert("end", os.path.basename(f_normalized))

    def mover_jpg_cima(self):
        sel = self.lst_jpgs.curselection()
        if not sel: return
        idx = sel[0]
        if idx == 0: return
        self.arquivos_jpg_lista[idx], self.arquivos_jpg_lista[idx - 1] = self.arquivos_jpg_lista[idx - 1], self.arquivos_jpg_lista[idx]
        text = self.lst_jpgs.get(idx)
        self.lst_jpgs.delete(idx)
        self.lst_jpgs.insert(idx - 1, text)
        self.lst_jpgs.selection_set(idx - 1)

    def mover_jpg_baixo(self):
        sel = self.lst_jpgs.curselection()
        if not sel: return
        idx = sel[0]
        if idx == self.lst_jpgs.size() - 1: return
        self.arquivos_jpg_lista[idx], self.arquivos_jpg_lista[idx + 1] = self.arquivos_jpg_lista[idx + 1], self.arquivos_jpg_lista[idx]
        text = self.lst_jpgs.get(idx)
        self.lst_jpgs.delete(idx)
        self.lst_jpgs.insert(idx + 1, text)
        self.lst_jpgs.selection_set(idx + 1)

    def remover_jpg_item(self):
        sel = self.lst_jpgs.curselection()
        if not sel: return
        idx = sel[0]
        self.lst_jpgs.delete(idx)
        self.arquivos_jpg_lista.pop(idx)

    def limpar_jpg_lista(self):
        self.lst_jpgs.delete(0, "end")
        self.arquivos_jpg_lista.clear()

    def iniciar_jpg_pdf(self):
        if not self.arquivos_jpg_lista:
            messagebox.showerror("Erro", "Sua lista está vazia!")
            return
        save_path = filedialog.asksaveasfilename(
            defaultextension=".pdf", filetypes=[("Documento PDF", "*.pdf")],
            initialfile=f"Imagens_Compiladas_Gree_{dt.now().strftime('%Y%m%d')}.pdf"
        )
        if save_path:
            save_path = os.path.normpath(save_path)
            orientation = self.cmb_orientacao.get()
            margin = self.cmb_margem.get()
            self.progress.set(0)
            self.lbl_porcentagem.configure(text="0%")
            self.escrever_log("Iniciando compilação...")
            thread = threading.Thread(target=self.worker_convert_jpg_to_pdf, args=(self.arquivos_jpg_lista, save_path, orientation, margin))
            thread.daemon = True
            thread.start()

    def worker_convert_jpg_to_pdf(self, images_list, output_path, orientation, margin):
        try:
            from PIL import Image
            self.queue.put(("log", f"Iniciando compilação..."))
            total_imgs = len(images_list)
            pil_images = []

            for i, img_path in enumerate(images_list):
                self.queue.put(("log", f"Processando ({i+1}/{total_imgs})..."))
                self.queue.put(("progress", int(((i + 1) / total_imgs) * 80)))

                img = Image.open(img_path)

                # TRATAMENTO MELHORADO PARA PNGs TRANSPARENTES E PALETAS
                if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
                    fundo_branco = Image.new("RGB", img.size, "white")
                    if img.mode == "P":
                        img = img.convert("RGBA")
                    fundo_branco.paste(img, mask=img.split()[-1])
                    img = fundo_branco
                elif img.mode != "RGB":
                    img = img.convert("RGB")

                # ==========================================
                # O SEGREDO DA RESOLUÇÃO: TELA A4 EM 300 DPI
                # ==========================================
                if orientation == "Paisagem (Horizontal)":
                    a4_w, a4_h = 3508, 2480
                else:
                    a4_w, a4_h = 2480, 3508

                # Ajustando as margens para a nova proporção de pixels gigantesca (300 DPI)
                if margin == "Margem fina": margin_px = 80
                elif margin == "Margem larga": margin_px = 250
                else: margin_px = 0

                target_w = a4_w - (2 * margin_px)
                target_h = a4_h - (2 * margin_px)
                img_w, img_h = img.size

                # Mantém o Lanczos, pois com um quadro de 3500px, a imagem não sofre compressão nociva
                ratio = min(target_w / img_w, target_h / img_h)
                new_w = int(img_w * ratio)
                new_h = int(img_h * ratio)

                resized_img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)

                # Cria a base em branco e foca no centro
                canvas = Image.new("RGB", (a4_w, a4_h), "white")
                offset_x = (a4_w - new_w) // 2
                offset_y = (a4_h - new_h) // 2
                canvas.paste(resized_img, (offset_x, offset_y))

                pil_images.append(canvas)

            if not pil_images:
                self.queue.put(("error", "Nenhuma imagem lida!"))
                return

            self.queue.put(("log", "Salvando PDF em HD (300 DPI)..."))
            self.queue.put(("progress", 90))

            # Salva apontando a densidade correta da nova tela
            pil_images[0].save(
                output_path,
                "PDF",
                resolution=300.0,
                save_all=True,
                append_images=pil_images[1:],
                quality=100,
                subsampling=0
            )

            self.queue.put(("progress", 100))
            self.queue.put(("success", f"PDF gerado em Alta Definição!\nSalvo em: {output_path}"))

        except Exception as e:
            self.queue.put(("error", f"Erro: {str(e)}"))

# ==========================================
# CÓDIGO PRINCIPAL: HUB LOGÍSTICA (Layout Modernizado)
# ==========================================
class HubLogistica(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("GREE - Hub de Ferramentas Logística")
        try:
            self.iconbitmap(CAMINHO_ICO)
        except Exception:
            try:
                from PIL import Image, ImageTk
                logo_icone = ImageTk.PhotoImage(Image.open(CAMINHO_LOGO))
                self.iconphoto(False, logo_icone)
            except Exception:
                pass

        self.cor_sidebar = "#0A2540"
        self.cor_ativa = "#F26522"
        self.cor_hover = "#123150"
        self.cinza_fundo = "#F4F6F9"

        largura_app = 1180
        altura_app = 780

        largura_tela = self.winfo_screenwidth()
        altura_tela = self.winfo_screenheight()

        pos_x = (largura_tela // 2) - (largura_app // 2)
        pos_y = (altura_tela // 2) - (altura_app // 2)

        self.geometry(f"{largura_app}x{altura_app}+{pos_x}+{pos_y}")

        self.sidebar = None
        self.content_frame = None
        self.ferramenta_atual = None

        self.construir_layout_principal()

        self.mostrar_ferramenta("CRM")

    def construir_layout_principal(self):
        self.sidebar = ctk.CTkFrame(self, width=260, corner_radius=0, fg_color=self.cor_sidebar)
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)

        try:
            logo_pil = Image.open(CAMINHO_LOGO)
            logo_img = ctk.CTkImage(light_image=logo_pil, dark_image=logo_pil, size=(170, 40))
            self.lbl_logo = ctk.CTkLabel(self.sidebar, image=logo_img, text="")
            self.lbl_logo.pack(pady=(40, 20))
        except Exception:
            self.lbl_logo = ctk.CTkLabel(
                self.sidebar,
                text="GREE LOGÍSTICA",
                font=("Segoe UI", 20, "bold"),
                text_color="white"
            )
            self.lbl_logo.pack(pady=(45, 20))

        divisoria = ctk.CTkFrame(self.sidebar, height=1, fg_color="#1C3F60")
        divisoria.pack(fill="x", padx=25, pady=(0, 25))

        font_menu = ("Segoe UI", 14, "bold")
        estilo_btn = {
            "fg_color": "transparent",
            "text_color": "#B0C4DE",
            "hover_color": self.cor_hover,
            "anchor": "w",
            "height": 45,
            "corner_radius": 8,
            "font": font_menu
        }

        self.btn_crm_side = ctk.CTkButton(
            self.sidebar, text="  📊   CRM Formatação", command=lambda: self.mostrar_ferramenta("CRM"), **estilo_btn)
        self.btn_crm_side.pack(fill="x", padx=15, pady=5)

        self.btn_cot_side = ctk.CTkButton(
            self.sidebar, text="  💰   Gerador de Cotações", command=lambda: self.mostrar_ferramenta("COTACAO"), **estilo_btn)
        self.btn_cot_side.pack(fill="x", padx=15, pady=5)

        self.btn_consolidador_side = ctk.CTkButton(
            self.sidebar, text="  📦   Consolidador PDFs", command=lambda: self.mostrar_ferramenta("CONSOLIDADOR"), **estilo_btn)
        self.btn_consolidador_side.pack(fill="x", padx=15, pady=5)

        texto_assinatura = (
            "Desenvolvido por:\n"
            "Waldir Neto & Vinicius Diogo\n\n"
            "\"Movendo o futuro da Gree\"\n"
            "推动格力的未来"
        )
        self.lbl_desenvolvedores = ctk.CTkLabel(
            self.sidebar,
            text=texto_assinatura,
            font=("Segoe UI", 11),
            text_color="#5D7A99",
            justify="center"
        )
        self.lbl_desenvolvedores.pack(side="bottom", pady=30)

        self.content_frame = ctk.CTkFrame(self, fg_color=self.cinza_fundo, corner_radius=0)
        self.content_frame.pack(side="right", fill="both", expand=True)

    def mostrar_ferramenta(self, nome):
        if self.ferramenta_atual:
            self.ferramenta_atual.destroy()

        for btn in [self.btn_crm_side, self.btn_cot_side, self.btn_consolidador_side]:
            btn.configure(fg_color="transparent", text_color="#B0C4DE", hover_color=self.cor_hover)

        if nome == "CRM":
            self.ferramenta_atual = CRMFormatter(self.content_frame)
            self.btn_crm_side.configure(fg_color=self.cor_ativa, text_color="white", hover_color="#D85A1E")

        elif nome == "COTACAO":
            self.ferramenta_atual = CotacaoApp(self.content_frame)
            self.btn_cot_side.configure(fg_color=self.cor_ativa, text_color="white", hover_color="#D85A1E")

        elif nome == "CONSOLIDADOR":
            self.ferramenta_atual = ConsolidadorApp(self.content_frame)
            self.btn_consolidador_side.configure(fg_color=self.cor_ativa, text_color="white", hover_color="#D85A1E")

        self.ferramenta_atual.pack(fill="both", expand=True)

if __name__ == "__main__":
    try:
        import ctypes
        myappid = 'gree.logistica.hub.v1'
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
    except Exception:
        pass

    app = HubLogistica()
    app.mainloop()